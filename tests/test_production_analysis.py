"""Production analysis report (2026-10-03): the owner's "Format Production
analysis" sheet, columns B..S + U..X, computed from Daily Entry punches.

Formulas (Sheet1, with the owner's 2026-10-06 rule for efficiency: what is not
the operator's mistake is CREDITED to him):
    U Planned qty          = H / G
    V Total actual qty     = I + J
    W Overall productivity = V x G / H
    X Operator efficiency  = (V x G + setting credit + downtime credited) / H
      setting credit    = the standard setting time where an actual one was typed
                          (blank standard = 90), else 0
      downtime credited = no power + breakdown + tool problem + no load + other work
      No operator is NOT credited.
"""
from datetime import date
import io

import pytest

from engine import production_analysis as pa
from engine.models import Actual


# --- the formulas ---------------------------------------------------------- #
def test_the_four_columns_follow_the_sheet():
    # G=5, H=330, I=40 OK, J=2 rejected, 90 setting credit, 30 min downtime credited.
    planned, total, prod, eff = pa.metrics(5, 330, 40, 2, 90, 30)
    assert planned == pytest.approx(66)                 # 330 / 5
    assert total == 42                                  # 40 + 2
    assert prod == pytest.approx(42 / 66)               # V / U = V x G / H
    assert eff == pytest.approx((42 * 5 + 90 + 30) / 330)


def _set(actual, std=90.0, **kw):
    return Actual(so_no="S", item_code="I", entry_date=date(2026, 10, 5), qty_produced=0,
                  cycle_time_min=5, shift_minutes=630, std_setup_min=std,
                  actual_setup_min=actual, **kw)


def test_setting_credit_is_the_standard_only_when_a_setup_was_typed():
    """Owner, 2026-10-06: a setup typed earns the STANDARD, slower or faster; a
    job run on in continuation (no actual typed) earns none; a blank standard
    earns the default 90."""
    assert pa.setting_credit(_set(120)) == 90            # slower: the extra 30 is his
    assert pa.setting_credit(_set(60)) == 90             # faster: still the standard
    assert pa.setting_credit(_set(0)) == 0               # continuation: no setup typed
    assert pa.setting_credit(_set(120, std=0)) == 90     # blank standard: the default
    assert pa.setting_credit(_set(120, std=0), default=75) == 75
    assert pa.setting_credit(_set(120, std=45)) == 45    # a typed standard wins


def test_downtime_credited_is_everything_but_no_operator():
    a = _set(0, no_power_min=10, no_operator_min=20, machine_breakdown_min=30,
             tool_problem_min=40, no_load_min=50, other_work_min=60)
    assert pa.downtime_credited(a) == 10 + 30 + 40 + 50 + 60     # no operator left out


def test_the_owners_worked_example():
    """630 available, 100 pcs x 4 min, standard setting 90 typed as 120 actual,
    no power 30, no operator 20: earned 400 + 90 + 30 = 520 -> 82.5%;
    productivity 400 / 630 = 63.5%."""
    a = _line("A", 100, 4.0, mins=630, setup=120, std_setup_min=90,
              no_power_min=30, no_operator_min=20)
    (row,) = pa.shift_rows([a], None, 2026, 10, WORKING)
    assert row["Setting credit in Min"] == 90
    assert row["Downtime credited in Min"] == 30
    assert row["No operator (in Min)"] == 20
    assert row["Standard minutes earned"] == 400
    assert row["Operator minutes earned"] == 520
    assert row["Operator efficiency"] == pytest.approx(82.5, abs=0.05)
    assert row["Overall Productivity"] == pytest.approx(63.5, abs=0.05)


def test_nothing_to_divide_by_is_blank_never_zero():
    assert pa.metrics(None, 600, 10, 0, 0, 0)[0] is None      # no cycle time
    assert pa.metrics(None, 600, 10, 0, 0, 0)[3] is None
    assert pa.metrics(5, 0, 10, 0, 0, 0)[2] is None           # no minutes entered
    assert pa.metrics(5, 0, 10, 0, 90, 40)[3] is None


def test_ok_qty_is_produced_minus_rejected():
    """The order book has always counted good = produced - rejected, so a punch's
    OK qty (I) is produced - rejected and Total actual (V) is produced."""
    a = Actual(so_no="S", item_code="I", entry_date=date(2026, 9, 1),
               qty_produced=50, qty_rejected=5, cycle_time_min=2, shift_minutes=600)
    row = pa.report_row(a)
    assert row["Actual OK Qty"] == 45
    assert row["Rejected qty"] == 5
    assert row["Total Actual Qty"] == 50


def test_new_fields_round_trip_and_legacy_rows_default():
    a = Actual(so_no="S", item_code="I", entry_date=date(2026, 9, 1), machine="CNC1",
               cycle_time_min=4.5, shift_minutes=330, std_setup_min=90)
    b = Actual.from_json(a.to_json())
    assert (b.machine, b.cycle_time_min, b.shift_minutes, b.std_setup_min) == ("CNC1", 4.5, 330, 90)
    legacy = {k: v for k, v in a.to_json().items()
              if k not in ("machine", "cycle_time_min", "shift_minutes", "std_setup_min")}
    c = Actual.from_json(legacy)
    assert (c.machine, c.cycle_time_min, c.shift_minutes, c.std_setup_min) == ("", None, 0.0, 0.0)


def test_report_columns_match_the_sheet_order():
    assert pa.REPORT_COLUMNS[0] == "Date" and pa.REPORT_COLUMNS[1] == "Machine"
    assert pa.REPORT_COLUMNS[-6:] == ("Planned qty", "Total Actual Qty",
                                      "Setting credit in Min", "Downtime credited in Min",
                                      "Standard minutes earned", "Operator minutes earned")
    assert "Rate" not in pa.REPORT_COLUMNS
    # Per-line percentages are misleading (owner, 2026-10-04): never on an entry row.
    assert "Operator efficiency" not in pa.REPORT_COLUMNS
    assert not any("Productivity" in c for c in pa.REPORT_COLUMNS)


# --- shift and month totals (owner, 2026-10-04) ----------------------------- #
WORKING = {"1st shift": 630, "2nd shift": 570}     # meal breaks taken out
def _line(item, qty, ct, mins=660, setup=0.0, op="Operator A", day=5, shift="1st shift",
          machine="CNC1", **kw):
    return Actual(so_no="S", item_code=item, entry_date=date(2026, 10, day), shift=shift,
                  operator=op, process="CNC FIRST SIDE", qty_produced=qty,
                  cycle_time_min=ct, shift_minutes=mins, actual_setup_min=setup,
                  machine=machine, **kw)


def test_a_full_shift_on_three_items_is_one_hundred_percent_not_three_small_ones():
    """The owner's example: set up A, run A, set up B, run B, set up C, run C,
    filling the 660-minute shift exactly, every setup on its standard. Each line
    on its own reads far below 100% because the floor types the whole shift on
    every line; the shift is 100%."""
    lines = [_line("A", 30, 4.0, setup=90, std_setup_min=90),     # 90 + 120 = 210
             _line("B", 20, 6.0, setup=90, std_setup_min=90),     # 90 + 120 = 210
             _line("C", 25, 2.4, setup=180, std_setup_min=180)]   # 180 + 60 = 240 -> 660
    for a in lines:
        (_, _, _, eff), _ = pa.actual_metrics(a)
        assert eff < 0.4                          # what a per-line column would show
    (row,) = pa.shift_rows(lines, None, 2026, 10)
    assert row["Entries"] == 3
    assert row["Minutes available in shift"] == 660       # counted ONCE, not 3 x 660
    assert row["Actual setting time in Min"] == 360
    assert row["Setting credit in Min"] == 360
    assert row["Standard minutes earned"] == 300
    assert row["Operator minutes earned"] == 660
    assert row["Operator efficiency"] == 100.0
    assert row["Overall Productivity"] == pytest.approx(300 / 660 * 100, abs=0.05)


def test_one_line_shift_matches_the_sheet_formula():
    a = _line("A", 42, 5, mins=330, setup=100, no_power_min=30)
    (_, _, prod, eff), _ = pa.actual_metrics(a)
    (row,) = pa.shift_rows([a], None, 2026, 10)
    assert row["Operator efficiency"] == pytest.approx(eff * 100, abs=0.05)
    assert row["Overall Productivity"] == pytest.approx(prod * 100, abs=0.05)


def test_shifts_are_split_by_day_shift_and_operator():
    acts = [_line("A", 10, 6), _line("B", 10, 6, op="Operator B"),
            _line("A", 10, 6, shift="2nd shift", mins=600), _line("A", 10, 6, day=6)]
    rows = pa.shift_rows(acts, None, 2026, 10)
    assert len(rows) == 4
    assert all(r["Entries"] == 1 for r in rows)


def test_the_month_adds_minutes_and_never_averages_percentages():
    """Shift 1: 600 earned in 600 minutes (100%). Shift 2, a half shift of 300
    minutes: 60 earned + 90 standard setting = 150 (50%; the setup took 300).
    Averaging the percentages gives 75%; the month is 750 / 900 = 83.3%."""
    acts = [_line("A", 100, 6.0, mins=600, day=5),                          # 600 / 600
            _line("A", 10, 6.0, mins=300, setup=300, std_setup_min=90, day=6)]  # 150 / 300
    (row,) = pa.operator_month_rows(acts, None, 2026, 10)
    assert row["Shifts worked"] == 2 and row["Shifts counted"] == 2
    assert row["Minutes available in shift"] == 900
    assert row["Standard minutes earned"] == 660
    assert row["Operator minutes earned"] == 750
    # 750 / 900 = 83.3%; averaging 100% and 50% would give 75%.
    assert row["Operator efficiency"] == pytest.approx(750 / 900 * 100, abs=0.05)
    assert row["Overall Productivity"] == pytest.approx(660 / 900 * 100, abs=0.05)


def test_a_shift_with_a_missing_cycle_time_has_no_figure_and_leaves_the_month():
    acts = [_line("A", 100, 6.0, mins=600, day=5),
            _line("A", 50, 6.0, mins=600, day=6), _line("B", 50, None, mins=600, day=6)]
    shifts = pa.shift_rows(acts, None, 2026, 10)
    day2 = next(r for r in shifts if r["Date"] == "06-10-2026")
    assert day2["Operator efficiency"] is None
    assert "no cycle time" in day2["Note"]
    (month,) = pa.operator_month_rows(acts, None, 2026, 10)
    assert month["Shifts worked"] == 2 and month["Shifts counted"] == 1
    assert month["Operator efficiency"] == 100.0
    assert "1 shift left out" in month["Note"]


def test_minutes_typed_once_per_line_are_counted_once():
    acts = [_line("A", 50, 6.0, mins=600), _line("B", 50, 6.0, mins=600)]
    (row,) = pa.shift_rows(acts, None, 2026, 10)
    assert row["Minutes available in shift"] == 600
    assert row["Operator efficiency"] == 100.0
    assert row["Note"] == ""


def test_month_rows_sort_best_first_and_no_figure_last():
    acts = [_line("A", 50, 6.0, mins=600, op="Low"), _line("A", 100, 6.0, mins=600, op="High"),
            _line("A", 100, None, mins=600, op="Blank")]
    assert [r["Operator"] for r in pa.operator_month_rows(acts, None, 2026, 10)] == \
        ["High", "Low", "Blank"]


def test_a_shift_never_counts_its_meal_break():
    """Lunch 13:00-13:30 / dinner 22:00-22:30 (owner, 2026-10-04): a line typed
    with the old full shift (660 / 600) counts 630 / 570; fewer minutes typed
    (a half shift) count as typed."""
    full = [_line("A", 105, 6.0, mins=660)]                       # 630 earned
    (row,) = pa.shift_rows(full, None, 2026, 10, WORKING)
    assert row["Minutes entered"] == 660
    assert row["Working minutes in shift (after break)"] == 630
    assert row["Minutes available in shift"] == 630
    assert row["Operator efficiency"] == 100.0
    assert "630 working minutes" in row["Note"]
    assert "smaller of 660 entered and 630 working minutes" in row["Working"]
    night = [_line("A", 95, 6.0, mins=600, shift="2nd shift")]    # 570 earned
    assert pa.shift_rows(night, None, 2026, 10, WORKING)[0]["Operator efficiency"] == 100.0
    half = [_line("A", 50, 6.0, mins=300)]
    (h,) = pa.shift_rows(half, None, 2026, 10, WORKING)
    assert h["Minutes available in shift"] == 300 and h["Operator efficiency"] == 100.0
    (e,) = pa.monthly_rows(full, None, 2026, 10, WORKING)
    assert e["Planned qty"] == 105                                # 630 / 6, not 660 / 6


def test_the_daily_entry_form_defaults_to_working_minutes():
    import api.main as m
    from engine.config import Config
    assert m._shift_minutes(Config(scheduler="new")) == WORKING


def test_the_daily_entry_list_no_longer_shows_the_calculated_columns():
    row = _line("A", 10, 6.0).as_row()
    for gone in ("Planned Qty", "Total Actual Qty", "Overall Productivity", "Operator Efficiency"):
        assert gone not in row
    assert row["Cycle Time (min)"] == 6.0 and row["Minutes Available"] == 660


# --- the API ---------------------------------------------------------------- #
pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from engine import book_store  # noqa: E402
from engine.models import Order  # noqa: E402
from tests.sample_workbook import build_sample_bytes, ITEM_A  # noqa: E402

DAY = "2026-10-03"   # the report's first day (REPORT_START), and never in the future


def _api():
    import importlib
    import api.main as m
    importlib.reload(m)
    book_store.save_masters_bytes(build_sample_bytes())
    book_store.add_orders([Order("SO1", ITEM_A, ITEM_A, 100, date(2026, 12, 1))])
    c = TestClient(m.app)
    c.post("/login", data={"username": "anvitech", "password": "1930rail"})
    return m, c


def _post(c, **kw):
    body = dict(so_no="SO1", item_code=ITEM_A, entry_date=DAY, shift="1st shift",
                process="BANDSAW", operator="Operator One", qty_produced=10)
    body.update(kw)
    return c.post("/actuals", json=body)


def test_save_snapshots_the_cycle_time_from_the_process_master():
    _, c = _api()
    r = _post(c, machine="BS1", shift_minutes=330, cycle_time_min=999)  # browser value ignored
    assert r.status_code == 200, r.text
    a = book_store.load_actuals()[-1]
    assert a.cycle_time_min == 3.0          # BANDSAW's cycle time in the sample master
    assert a.machine == "BS1" and a.shift_minutes == 330


def test_standard_setting_time_defaults_by_machine_kind():
    _, c = _api()
    assert _post(c, machine="BS1", qty_produced=5).status_code == 200
    assert book_store.load_actuals()[-1].std_setup_min == 0.0        # manual station
    assert _post(c, process="CNC OS", machine="CNC1", qty_produced=5).status_code == 200
    assert book_store.load_actuals()[-1].std_setup_min == 90.0       # CNC
    assert _post(c, process="CNC OS", machine="CNC1", qty_produced=0,
                 std_setup_min=120).status_code == 200
    assert book_store.load_actuals()[-1].std_setup_min == 120.0      # typed value kept


def test_an_unknown_machine_is_refused():
    _, c = _api()
    r = _post(c, machine="CNC77")
    assert r.status_code == 400 and "machine list" in r.text


def test_items_feeds_the_form_its_defaults():
    _, c = _api()
    body = c.get("/items").json()
    assert body["shift_minutes"] == {"1st shift": 660, "2nd shift": 600}
    info = body["items"][ITEM_A]["process_info"]
    assert info["BANDSAW"] == {"cycle_time": 3.0, "machines": ["BS1"]}
    assert info["CNC OS"]["machines"] == ["CNC1", "CNC2"]
    setups = {m["id"]: m["std_setup_min"] for m in body["machines"]}
    assert setups["CNC1"] == 90 and setups["BS1"] == 0


def test_monthly_report_json_and_excel():
    _, c = _api()
    _post(c, machine="BS1", shift_minutes=330, qty_produced=12, qty_rejected=2,
          actual_setup_min=0, no_power_min=30)
    _post(c, machine="BS1", shift_minutes=330, qty_produced=8)
    r = c.get("/production-analysis", params={"year": 2026, "month": 10})
    assert r.status_code == 200
    body = r.json()
    e1, e2 = body["entries"]["rows"]
    assert e1["Machine"] == "BS1" and e1["Cycle time in Min"] == 3.0
    assert e1["Planned qty"] == 110                          # 330 / 3
    assert e1["Total Actual Qty"] == 12                      # 10 OK + 2 rejected
    assert e1["Standard minutes earned"] == 36               # 12 x 3
    (shift,) = body["shifts"]["rows"]                        # both lines, one shift
    assert shift["Entries"] == 2 and shift["Minutes available in shift"] == 330
    assert shift["Standard minutes earned"] == 60            # (12 + 8) x 3
    assert shift["Downtime credited in Min"] == 30           # the 30 min no power
    assert shift["Operator minutes earned"] == 90            # 60 + 30
    assert shift["Operator efficiency"] == pytest.approx(90 / 330 * 100, abs=0.05)
    (op,) = body["operators"]["rows"]
    assert op["Operator"] == "Operator One"
    assert op["Operator efficiency"] == shift["Operator efficiency"]
    empty = c.get("/production-analysis", params={"year": 2026, "month": 9}).json()
    assert empty["entries"]["rows"] == [] and empty["operators"]["rows"] == []

    x = c.get("/production-analysis.xlsx", params={"year": 2026, "month": 10})
    assert x.status_code == 200
    from openpyxl import load_workbook
    from engine import production_analysis_xlsx as pax
    values = load_workbook(io.BytesIO(x.content), data_only=True)
    formulas = load_workbook(io.BytesIO(x.content))
    assert values.sheetnames == ["Operator efficiency", "Shift-wise", "Every entry",
                                 "How it is calculated"]
    month_col = 2 + pa.MONTH_COLUMNS.index("Operator efficiency")
    assert values["Operator efficiency"].cell(row=7, column=month_col).value == \
        pytest.approx(90 / 330, abs=0.001)
    # Transparent: every calculated cell is a formula over the cells it comes from.
    f = formulas["Operator efficiency"].cell(row=7, column=month_col).value
    assert f.startswith("=IF(") and "/" in f
    sh = formulas["Shift-wise"]
    col = lambda name: 2 + pa.SHIFT_COLUMNS.index(name)
    assert sh.cell(row=7, column=col("Minutes entered")).value.startswith(
        "=MAX('Every entry'!")
    assert sh.cell(row=7, column=col("Minutes available in shift")).value.startswith(
        "=IF(ISNUMBER(")
    assert sh.cell(row=7, column=col("Standard minutes earned")).value.startswith(
        "=SUM('Every entry'!")
    en = formulas["Every entry"]
    assert en["B6"].value == "Date" and en["C6"].value == "Machine"
    earned = en.cell(row=7, column=2 + pa.REPORT_COLUMNS.index("Standard minutes earned"))
    assert earned.value.startswith("=IF(ISNUMBER(")
    working = values["Shift-wise"].cell(row=7, column=col("Working")).value
    assert working.startswith("Standard minutes earned = 12 x 3 + 8 x 3 = 60")


def test_every_excel_formula_gives_the_value_the_preview_shows():
    """Evaluate every formula in the download independently (pycel) and check it
    gives the stored value, which is the value the preview shows. Covers a
    multi-line shift, a second operator, and a shift left out for a missing
    cycle time."""
    pycel = pytest.importorskip("pycel")
    import pycel.excellib as xl
    if not hasattr(xl, "counta"):           # pycel lacks COUNTA: Excel's meaning
        def counta(*args):
            flat = []
            for a in args:
                flat.extend(v for r in a for v in r) if isinstance(a, tuple) else flat.append(a)
            return sum(1 for v in flat if v not in (None, ""))
        xl.counta = counta
    from engine import production_analysis_xlsx as pax
    import math
    # Setting credits of every kind: a typed standard (A, slower than it), a blank
    # standard (B, the default 90), a faster setup (C); no operator and other work
    # on the same lines, so the credited / not-credited split is checked too.
    acts = [_line("A", 30, 4.0, setup=120, std_setup_min=90, no_operator_min=25),
            _line("B", 20, 6.0, setup=90, other_work_min=15),
            _line("C", 25, 2.4, setup=60, std_setup_min=180),
            _line("A", 50, 6.0, mins=600, op="Operator B", shift="2nd shift"),
            _line("A", 10, 6.0, day=6), _line("D", 5, None, day=6),
            _line("A", 10, 6.0, op="Operator B", day=6, no_power_min=60)]
    wk = WORKING
    acts.append(_line("A", 20, 6.0, mins=300, op="Operator C"))     # under the cap
    tables = {"operators": {"rows": pa.operator_month_rows(acts, None, 2026, 10, wk)},
              "shifts": {"rows": pa.shift_rows(acts, None, 2026, 10, wk)},
              "entries": {"rows": pa.monthly_rows(acts, None, 2026, 10, wk)}}
    assert any(r["Minutes entered"] == 660 and r["Minutes available in shift"] == 630
               for r in tables["shifts"]["rows"])                     # the cap is exercised
    import tempfile, os
    data = pax.build(2026, 10, tables)
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "r.xlsx")
        open(path, "wb").write(data)
        wf, wv = load_workbook(path), load_workbook(path, data_only=True)
        # Re-saving through openpyxl drops every stored value, so pycel has to
        # RECALCULATE each formula from the inputs; given the original file it
        # would just hand back the stored values and prove nothing.
        bare = os.path.join(d, "bare.xlsx")
        wf.save(bare)
        xc = pycel.ExcelCompiler(filename=bare)
        n = 0
        for name in wf.sheetnames:
            for row in wf[name].iter_rows(min_row=7):
                for cell in row:
                    if cell.data_type == "f":
                        n += 1
                        got = xc.evaluate(f"'{name}'!{cell.coordinate}")
                        want = wv[name][cell.coordinate].value
                        if isinstance(want, (int, float)):
                            # The stored value is the preview's, rounded to 0.1 (0.1% for
                            # a percentage); the formula's exact result must round to it.
                            head = wf[name].cell(row=6, column=cell.column).value
                            tol = 0.0005 if head in pa.PCT_COLUMNS else 0.05
                            assert abs(got - want) <= tol + 1e-9, (name, cell.coordinate, got, want)
                        else:
                            assert got == want, (name, cell.coordinate, got, want)
        assert n > 50
        # Text is never a formula: the only formulas are the calculated cells.
        how = wf["How it is calculated"]
        assert not any(c.data_type == "f" for row in how.iter_rows() for c in row)
        labels = {how.cell(row=r, column=2).value: how.cell(row=r, column=3).value
                  for r in range(3, how.max_row + 1)}
        assert labels["Every entry: Total Actual Qty"] == "= Actual OK Qty + Rejected qty"


def test_entries_before_the_report_start_are_left_out():
    """The form first had the report's fields on 03-10-2026; a 02-10 line has no
    minutes or machine and would only show as a shift 'left out'."""
    acts = [_line("A", 10, 6.0, day=2), _line("A", 10, 6.0, day=3)]
    assert [r["Date"] for r in pa.monthly_rows(acts, None, 2026, 10)] == ["03-10-2026"]
    assert len(pa.shift_rows(acts, None, 2026, 10)) == 1
    (month,) = pa.operator_month_rows(acts, None, 2026, 10)
    assert month["Shifts worked"] == 1
    assert pa.monthly_rows([_line("A", 10, 6.0, day=30)], None, 2026, 9) == []


def test_report_fields_never_reach_the_plan():
    """The new fields are report-only: two punches that differ ONLY in them give
    the planner identical remaining quantities."""
    from engine import orderbook
    base = dict(so_no="SO1", item_code=ITEM_A, entry_date=date(2026, 9, 1),
                process="BANDSAW", operator="Operator One", qty_produced=10)
    plain = Actual(**base)
    rich = Actual(**base, machine="BS1", cycle_time_min=3, shift_minutes=330, std_setup_min=0)
    orders = {("SO1", ITEM_A): Order("SO1", ITEM_A, ITEM_A, 100, date(2026, 12, 1))}
    lines = lambda acts: [(l.so_no, l.item_code, l.qty, l.process_qty)
                          for l in orderbook.active_so_lines(orders, acts)]
    assert lines([plain]) == lines([rich])


# --- the cycle time comes from the CURRENT master (owner, 2026-10-04) ----------- #
def _master(item="I", process="CNC FIRST SIDE", ct=45.0):
    from types import SimpleNamespace as NS
    from engine.models import Process, Routing
    proc = Process(seq=1, name=process, cycle_time=ct, total_time=ct,
                   suggested_machine="CNC1", allotted_machine="CNC1")
    return NS(routings={item: Routing(item_code=item, description="", customer="",
                                      rm_type="", moq=None, processes=[proc])})


def test_the_current_master_wins_over_the_value_saved_on_the_punch():
    """A punch saved while the master carried a padded 54 reads the master's 45 now:
    the report always uses the master on file, never a stale snapshot."""
    a = _line("I", 10, 54.0, mins=630)
    row = pa.report_row(a, _master(ct=45.0), WORKING)
    assert row["Cycle time in Min"] == 45
    assert row["Standard minutes earned"] == 450              # 10 x 45, not 10 x 54
    (shift,) = pa.shift_rows([a], _master(ct=45.0), 2026, 10, WORKING)
    assert shift["Standard minutes earned"] == 450
    (month,) = pa.operator_month_rows([a], _master(ct=45.0), 2026, 10, WORKING)
    assert month["Standard minutes earned"] == 450


def test_a_step_no_longer_in_the_master_falls_back_to_the_saved_value():
    a = _line("I", 10, 54.0, mins=630)
    for master in (_master(item="OTHER"), _master(process="RENAMED STEP"),
                   _master(ct=None)):
        assert pa.report_row(a, master, WORKING)["Cycle time in Min"] == 54
