"""Add New Orders: storage, endpoints, role gating (2026-09-08 spec)."""
import importlib
import io
import time
from datetime import date, timedelta

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient

from engine import book_store, loaders, orderbook
from engine.config import Config
from engine.models import Order, PlanRun, fmt_date
from tests.new_sample_workbook import build_new_sample_bytes, ITEM_A


def _rows(table):
    """``to_table()``'s rows are plain VALUE LISTS in ``table["columns"]``
    order, not dicts — zip them back into dicts for readable assertions."""
    return [dict(zip(table["columns"], row)) for row in table["rows"]]


def test_drafts_round_trip():
    book_store.save_new_order_drafts([{"so_no": "NEW-1", "item_code": "X",
                                       "item_name": "x", "qty": 10}])
    assert book_store.load_new_order_drafts() == [
        {"so_no": "NEW-1", "item_code": "X", "item_name": "x", "qty": 10}]


def test_drafts_default_to_empty():
    book_store.save_new_order_drafts([])
    assert book_store.load_new_order_drafts() == []


def test_queue_round_trips_and_clears():
    """2026-09-11: the queue is a list of arrival GROUPS, not a flat list of
    pairs — each group is the set of lines one accept added together, so
    _plan can replan a group POOLED (see the module docstring)."""
    book_store.save_new_order_queue([[("NEW-1", "X")], [("NEW-2", "Y")]])
    assert book_store.load_new_order_queue() == [[["NEW-1", "X"]], [["NEW-2", "Y"]]]
    book_store.clear_new_order_queue()
    assert book_store.load_new_order_queue() == []


def test_a_multi_line_group_round_trips_as_one_group():
    """A group with more than one line (one accept, several lines) stays one
    group, in the order it was written — not flattened or reordered."""
    book_store.save_new_order_queue([[("NEW-1", "X"), ("NEW-2", "Y")]])
    assert book_store.load_new_order_queue() == [[["NEW-1", "X"], ["NEW-2", "Y"]]]


def test_a_queue_written_in_the_old_flat_shape_still_loads():
    """Pre-2026-09-11 data: a flat list of pairs, one entry per (so_no,
    item_code), with no record of which lines arrived together. Loading it
    now must not crash, and must not silently invent clubbing that was never
    recorded — each old entry becomes its own one-line group."""
    old_shape = [["NEW-1", "X"], ["NEW-2", "Y"], ["NEW-3", "Z"]]
    book_store.get_store().kv_set(
        book_store.NEW_ORDER_QUEUE_KEY, book_store.json.dumps(old_shape))
    assert book_store.load_new_order_queue() == [
        [["NEW-1", "X"]], [["NEW-2", "Y"]], [["NEW-3", "Z"]]]


# --- validation ---

def _masters():
    _so, masters = loaders.load_all(io.BytesIO(build_new_sample_bytes()))
    return masters


def _an_item(masters):
    return sorted(masters.routings)[0]


def test_a_good_line_validates():
    m = _masters()
    assert orderbook.validate_new_order_line("NEW-1", _an_item(m), 10, {}, {}, m) is None


def test_a_blank_so_number_is_refused():
    m = _masters()
    assert "SO number" in orderbook.validate_new_order_line("  ", _an_item(m), 10, {}, {}, m)


def test_a_zero_or_negative_quantity_is_refused():
    m = _masters()
    for bad in (0, -5):
        assert "quantity" in orderbook.validate_new_order_line("NEW-1", _an_item(m),
                                                               bad, {}, {}, m).lower()


def test_an_unknown_item_is_refused():
    m = _masters()
    assert "routing" in orderbook.validate_new_order_line("NEW-1", "GHOST", 10,
                                                          {}, {}, m).lower()


def test_a_duplicate_names_where_it_already_is():
    m = _masters()
    item = _an_item(m)
    existing = Order(so_no="NEW-1", item_code=item, item_name="x", ordered_qty=300,
                     delivery_date=date(2025, 3, 20))
    msg = orderbook.validate_new_order_line("new-1", item, 10,
                                            {existing.key: existing}, {}, m)
    assert "already in the book" in msg
    assert "300" in msg


def test_a_duplicate_of_a_COMPLETED_order_is_also_refused():
    m = _masters()
    item = _an_item(m)
    done = Order(so_no="NEW-1", item_code=item, item_name="x", ordered_qty=300,
                 delivery_date=date(2025, 3, 20), completed=True)
    assert "already in the book" in orderbook.validate_new_order_line(
        "NEW-1", item, 10, {}, {done.key: done}, m)


# --- occupancy is new-engine-only (2026-09-08, Task 10 requirement) --- #

def test_occupancy_is_rejected_outside_the_new_engine():
    """``rule6_allocate.run`` (classic) and ``flow_scheduler.run`` (flow) both
    declare ``**kw`` and would silently discard ``occupancy`` — a two-stage
    quote on either of those would look like it protects the existing plan
    while doing nothing. ``run_forward`` must refuse loudly instead."""
    from engine.pipeline import OccupancyRequiresNewEngineError, run_forward

    masters = _masters()
    cfg = Config(plan_start_date=date(2025, 3, 1), scheduler="classic")
    plan_run = PlanRun(so_lines=[])
    with pytest.raises(OccupancyRequiresNewEngineError) as exc:
        run_forward(plan_run, cfg, masters, occupancy={"machine": {}, "operator": {}})
    assert "classic" in str(exc.value)


# --- Task 10: draft endpoints and the quote endpoint --- #

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _api_with_new_engine(monkeypatch):
    """These endpoints plan through ``run_forward(occupancy=...)``, which only the
    new engine honours (see the ``OccupancyRequiresNewEngineError`` test above) —
    force ``DEFAULT_SCHEDULER=new`` and confirm the pin actually took. Without
    this, every test below would either fail loudly for the wrong reason or
    (with the env unset, as in the rest of the suite) run on the classic engine
    and prove nothing about the feature it claims to test."""
    monkeypatch.setenv("DEFAULT_SCHEDULER", "new")
    import api.main as m
    importlib.reload(m)
    resolved = m._resolve_config(m._load_plan_config())
    assert resolved.scheduler == "new", (
        f"DEFAULT_SCHEDULER=new did not take: resolved scheduler is "
        f"{resolved.scheduler!r}")
    return m


@pytest.fixture
def _api_module(monkeypatch):
    return _api_with_new_engine(monkeypatch)


@pytest.fixture
def admin_client(_api_module):
    c = TestClient(_api_module.app)
    r = c.post("/login", data={"username": "anvitech", "password": "1930rail"})
    assert r.status_code in (200, 303), r.text
    return c


@pytest.fixture
def user_client(_api_module):
    c = TestClient(_api_module.app)
    r = c.post("/login", data={"username": "anvitech_user", "password": "anvitech12345678"})
    assert r.status_code in (200, 303), r.text
    return c


@pytest.fixture
def uploaded_masters(admin_client):
    """Upload the new-engine sample workbook (admin only) and return one item
    code it defines, for a draft line's item_code."""
    r = admin_client.post("/upload", files={
        "file": ("new.xlsx", build_new_sample_bytes(), XLSX_MIME)})
    assert r.status_code == 200, r.text
    return ITEM_A


@pytest.fixture
def add_new_order(admin_client):
    """PUT a single-line draft, quote it, accept it, and return the quoted
    completion date (an ISO date string) — the whole "type it, quote it,
    accept it" flow a director drives by hand, in one call, for the tests
    that only care about what ends up in the book (Task 11/12, 2026-09-08)."""
    def _add(so_no, item_code, qty):
        r = admin_client.put("/new-orders/drafts",
                             json={"drafts": [{"so_no": so_no, "item_code": item_code,
                                               "qty": qty}]})
        assert r.status_code == 200, r.text
        quote = admin_client.post("/new-orders/quote").json()
        added = admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
        assert added.status_code == 200, added.text
        return quote["lines"][0]["completion"]
    return _add


@pytest.fixture
def upload_sample(admin_client):
    """Re-upload the new-engine sample workbook — a fresh upload, not the
    original one `uploaded_masters` already did."""
    def _upload():
        r = admin_client.post("/upload", files={
            "file": ("new2.xlsx", build_new_sample_bytes(), XLSX_MIME)})
        assert r.status_code == 200, r.text
        return r
    return _upload


def test_drafts_are_admin_only(user_client):
    assert user_client.get("/new-orders/drafts").status_code == 403
    assert user_client.put("/new-orders/drafts", json={"drafts": []}).status_code == 403
    assert user_client.post("/new-orders/quote").status_code == 403


def test_saving_and_reading_back_drafts(admin_client, uploaded_masters):
    item = uploaded_masters
    r = admin_client.put("/new-orders/drafts",
                         json={"drafts": [{"so_no": "NEW-1", "item_code": item,
                                           "qty": 10}]})
    assert r.status_code == 200
    got = admin_client.get("/new-orders/drafts").json()
    assert got["drafts"][0]["so_no"] == "NEW-1"
    assert got["drafts"][0]["item_name"]
    assert any(i["item_code"] == item for i in got["items"])


def test_a_bad_draft_line_is_refused_with_a_message(admin_client, uploaded_masters):
    r = admin_client.put("/new-orders/drafts",
                         json={"drafts": [{"so_no": "", "item_code": uploaded_masters,
                                           "qty": 10}]})
    assert r.status_code == 400
    assert "SO number" in r.json()["detail"]


def test_a_quote_returns_a_date_and_confirms_nothing_moved(admin_client,
                                                           uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 10}]})
    r = admin_client.post("/new-orders/quote")
    assert r.status_code == 200
    body = r.json()
    assert body["verified"] is True
    assert body["moved"] == []
    assert body["lines"][0]["completion"]
    assert body["existing_count"] >= 0
    assert body["stamp"]


def test_a_refused_quote_always_carries_a_plain_reason(admin_client, uploaded_masters,
                                                        monkeypatch):
    """2026-09-11 review, smaller finding: a refusal caused by
    `structural_violations` alone left `moved` empty and the response never
    even included `violations`, so a director reading `verified: false` had
    no idea why. Stubbed directly since `verify_unmoved` can never actually
    see a move through this endpoint's own natural flow (it recomputes
    `existing_expected` from the same schedule it hands to `quote()` as
    `existing_entries`)."""
    from engine.quote import QuoteResult
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 10}]})
    monkeypatch.setattr(
        "engine.quote.quote",
        lambda *a, **kw: QuoteResult(
            lines=[], moved=[], violations=["machine CNC1 is double-booked"],
            verified=False, existing_count=0, order=[]))
    r = admin_client.post("/new-orders/quote")
    assert r.status_code == 200
    body = r.json()
    assert body["verified"] is False
    assert body["violations"] == ["machine CNC1 is double-booked"]
    assert body["reason"], "a refusal must always carry a reason"
    assert "double-booked" in body["reason"]


def test_a_refused_quote_reason_counts_moved_orders_too(admin_client, uploaded_masters,
                                                        monkeypatch):
    """The other half of the same fix: a refusal caused by `moved` (existing
    orders that would move) still has to say so in plain English."""
    from engine.quote import QuoteResult
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 10}]})
    monkeypatch.setattr(
        "engine.quote.quote",
        lambda *a, **kw: QuoteResult(
            lines=[], moved=[{"so_no": "X", "item_code": "Y",
                             "before": date(2025, 1, 1), "after": date(2025, 1, 2),
                             "days": 1}],
            violations=[], verified=False, existing_count=0, order=[]))
    r = admin_client.post("/new-orders/quote")
    body = r.json()
    assert body["verified"] is False
    assert "1 existing order" in body["reason"]


def test_quoting_with_no_drafts_is_a_clear_400(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts", json={"drafts": []})
    r = admin_client.post("/new-orders/quote")
    assert r.status_code == 400
    assert "no new orders" in r.json()["detail"].lower()


def test_the_quote_honours_machine_downtime(admin_client, uploaded_masters):
    """2026-09-08 review, Finding 1 (CRITICAL): neither `/new-orders/quote` nor
    `/new-orders/add` passed `reserved=` into `quote_mod.quote`, so the quote
    could promise a date inside a machine's maintenance window while the
    plan's own two-stage pass (which already honours downtime) computed a
    later one for the very same order once it was added. CNC1 is Item A's
    Allotted machine for its first step (new_sample_workbook.py) with no
    Suggested fallback, so taking it out of service for a few days can only
    delay the order, never make it unschedulable — the discriminating case
    the review asked for. The quoted date must equal what `/run` then reports
    for the very order that was just added."""
    today = date.today()
    book_store.save_machine_downtime({
        "machine": "CNC1", "from_date": today.isoformat(),
        "to_date": (today + timedelta(days=5)).isoformat(), "reason": "test"})
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    quote = admin_client.post("/new-orders/quote").json()
    quoted = quote["lines"][0]["completion"]
    assert quoted, "the order should still be schedulable around a 5-day break"
    r = admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
    assert r.status_code == 200, r.text
    body = admin_client.post("/run").json()
    actual = body["expected_end"][f"NEW-1\x1f{uploaded_masters}"]
    assert actual == quoted, (
        f"the quote promised {quoted} but the plan's own completion is "
        f"{actual}: the quote ignored machine downtime")


# --- Task 10 review fix: the intra-payload dedup + no-partial-persistence, --- #
# --- both correct but untested before (2026-09-08) --- #

def test_a_duplicate_line_within_the_same_payload_is_refused(admin_client,
                                                              uploaded_masters):
    r = admin_client.put("/new-orders/drafts",
                         json={"drafts": [
                             {"so_no": "NEW-1", "item_code": uploaded_masters, "qty": 10},
                             {"so_no": "NEW-1", "item_code": uploaded_masters, "qty": 5}]})
    assert r.status_code == 400
    assert "twice" in r.json()["detail"]


def test_a_bad_line_saves_no_line_from_the_same_payload(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts", json={"drafts": []})   # start clean
    r = admin_client.put("/new-orders/drafts",
                         json={"drafts": [
                             {"so_no": "NEW-1", "item_code": uploaded_masters, "qty": 10},
                             {"so_no": "", "item_code": uploaded_masters, "qty": 5}]})
    assert r.status_code == 400
    # The first (valid) line must NOT have been persisted just because it came
    # before the bad one in the same payload — the whole PUT is all-or-nothing.
    assert admin_client.get("/new-orders/drafts").json()["drafts"] == []


# --- Task 12: accepting a quote, and clearing the queue --- #

def test_adding_saves_the_orders_with_the_quoted_delivery_date(admin_client,
                                                               uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    quote = admin_client.post("/new-orders/quote").json()
    r = admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
    assert r.status_code == 200
    # GET /orders returns the same to_table() shape as POST /run's "orders" key:
    # {"columns": [...], "rows": [...]} — not a bare list of dicts — and the
    # delivery date column is "SO Delivery Date", rendered DD-MM-YYYY (fmt_date),
    # not the ISO string the quote returns.
    rows = {(o["SO No"], o["Item Code"]): o
            for o in _rows(admin_client.get("/orders").json()["orders"])}
    row = rows[("NEW-1", uploaded_masters)]
    quoted_date = date.fromisoformat(quote["lines"][0]["completion"])
    assert row["SO Delivery Date"] == fmt_date(quoted_date)


def test_adding_clears_the_drafts(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    quote = admin_client.post("/new-orders/quote").json()
    admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
    assert admin_client.get("/new-orders/drafts").json()["drafts"] == []


def test_a_stale_quote_is_refused_and_says_to_requote(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    admin_client.post("/new-orders/quote")
    r = admin_client.post("/new-orders/add", json={"stamp": "not-the-real-stamp"})
    assert r.status_code == 409
    assert "Finish and Optimize" in r.json()["detail"]


def test_add_is_refused_when_the_quote_could_not_be_confirmed(admin_client,
                                                              uploaded_masters,
                                                              monkeypatch):
    """2026-09-08 review, minor item: the ``not res.verified`` branch had no
    test. ``quote_mod.quote`` is stubbed directly — under normal operation
    the endpoint recomputes ``existing_expected`` from the SAME schedule it
    passes as ``existing_entries``, so ``verify_unmoved`` can never actually
    see a move by construction; this exercises the branch in isolation."""
    from engine.quote import QuoteResult
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    quote = admin_client.post("/new-orders/quote").json()
    monkeypatch.setattr(
        "engine.quote.quote",
        lambda *a, **kw: QuoteResult(
            lines=[], moved=[{"so_no": "X", "item_code": "Y",
                             "before": date(2025, 1, 1), "after": date(2025, 1, 2),
                             "days": 1}],
            violations=[], verified=False, existing_count=0, order=[]))
    r = admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
    assert r.status_code == 409
    assert "could not be confirmed" in r.json()["detail"]


def test_add_surfaces_the_error_when_a_line_could_not_be_scheduled(admin_client,
                                                                   uploaded_masters,
                                                                   monkeypatch):
    """2026-09-08 review, minor item: the ``completion is None`` branch had
    no test. Stubbed the same way as the verified=False test above."""
    from engine.quote import QuoteResult
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    quote = admin_client.post("/new-orders/quote").json()
    monkeypatch.setattr(
        "engine.quote.quote",
        lambda *a, **kw: QuoteResult(
            lines=[{"so_no": "NEW-1", "item_code": uploaded_masters, "item_name": "x",
                    "qty": 25, "target_date": None, "completion": None,
                    "error": "could not be scheduled: no machine with a qualified "
                             "operator is available for one of its steps"}],
            moved=[], violations=[], verified=True, existing_count=0, order=[]))
    r = admin_client.post("/new-orders/add", json={"stamp": quote["stamp"]})
    assert r.status_code == 400
    assert "could not be scheduled" in r.json()["detail"]


def test_adding_is_admin_only(user_client):
    assert user_client.post("/new-orders/add", json={"stamp": "x"}).status_code == 403


def test_done_entering_clears_the_queue_when_a_contest_actually_starts(
        admin_client, uploaded_masters, add_new_order, monkeypatch):
    """Clearing the queue is only correct when Done actually launches a
    re-optimization: `_try_start_auto` returning True means a background
    contest was launched (see test_auto_optimize.py's own pattern for
    stubbing `_start_optimize` rather than waiting on a real search)."""
    add_new_order("NEW-1", uploaded_masters, 25)
    assert book_store.load_new_order_queue()
    import api.main as m
    monkeypatch.setenv("AUTO_OPTIMIZE", "1")
    monkeypatch.setattr(
        m, "_start_optimize",
        lambda budget_evals, label, background=True, auto=False: None)
    r = admin_client.post("/optimize/done")
    assert r.json()["started"] is True
    assert book_store.load_new_order_queue() == []


def test_done_entering_does_not_clear_the_queue_when_it_skips(
        admin_client, uploaded_masters, add_new_order):
    """AUTO_OPTIMIZE is off by default in the test environment, so Done always
    skips here — clearing the queue on a skip would drop the arrival
    positions with no re-optimization ever having happened (2026-09-08
    review finding: the minor item about the early clear)."""
    add_new_order("NEW-1", uploaded_masters, 25)
    r = admin_client.post("/optimize/done")
    assert r.json()["started"] is False
    assert book_store.load_new_order_queue() == [[["NEW-1", uploaded_masters]]]


def test_an_upload_clears_the_queue(admin_client, uploaded_masters, add_new_order,
                                    upload_sample):
    add_new_order("NEW-1", uploaded_masters, 25)
    upload_sample()
    assert book_store.load_new_order_queue() == []


def test_applying_an_optimization_clears_the_queue(admin_client, uploaded_masters,
                                                   add_new_order):
    """2026-09-08 review finding 4: `_optimize_apply()` is the canonical
    "full optimization" of the three clearing sites (Done entering, an
    applied deep search, an upload) and was untested — deleting the clear
    there left the whole suite green."""
    add_new_order("NEW-1", uploaded_masters, 25)
    assert book_store.load_new_order_queue()
    import api.main as m
    st = m._start_optimize(budget_evals=15, label="quick", background=False)
    assert st["state"] == "done", st
    m._optimize_apply()
    assert book_store.load_new_order_queue() == []


def test_delay_report_includes_a_queued_new_order(admin_client, uploaded_masters,
                                                   add_new_order):
    """2026-09-11 review, I4: `_plan` cached STAGE 1's `so_lines` only under
    `_PLAN_CACHE["artifacts"]`, while `plan_run.schedule` in that same cache
    entry already carried BOTH stages merged (stage 2 folds the queued order
    in). `_plan_run_for_report` hands those artifacts straight to
    `build_delay_report`, which enumerates `so_lines` to decide which orders
    to explain — so a queued new order had real rows in the schedule but ZERO
    rows in the delay report, the silent-omission class CLAUDE.md names
    explicitly. Proven directly against `build_delay_report` (what the
    endpoint calls), on the SAME schedule, so the before/after counts below
    are both real and both come from the one run: before = the stage-1-only
    line list the old code cached, after = the merged list this fix caches."""
    import api.main as m
    from engine import delay_report as dr

    add_new_order("NEW-1", uploaded_masters, 25)
    assert book_store.load_new_order_queue(), "the new order must still be queued"

    m._plan(m._load_plan_config())
    art = m._PLAN_CACHE["artifacts"]
    plan_run, all_lines, masters, cfg = (art["plan_run"], art["so_lines"],
                                         art["masters"], art["config"])

    # The fix, checked directly: the cached list already names the queued order.
    assert any(l.so_no == "NEW-1" for l in all_lines), (
        "the artifacts' so_lines must include the queued new order (I4 fix)")
    assert any(e.so_refs and "NEW-1" in e.so_refs for e in plan_run.schedule), (
        "sanity: the merged schedule really does contain the new order's work")

    def _rows_for_new_order(lines):
        report = dr.build_delay_report(plan_run.schedule, lines,
                                        plan_run.batches_prioritized, cfg, masters,
                                        book_store.load_machine_downtime())
        return [r for r in report["detail"] if r["SO No"] == "NEW-1"]

    before_rows = _rows_for_new_order([l for l in all_lines if l.so_no != "NEW-1"])
    after_rows = _rows_for_new_order(all_lines)
    print(f"I4: delay report detail rows for NEW-1 — before={len(before_rows)} "
         f"after={len(after_rows)}")
    assert before_rows == [], "sanity: the stage-1-only list must not name NEW-1"
    assert after_rows, "the queued new order must have rows in the delay report"

    # And the real endpoint (which reads these same cached artifacts) agrees.
    pytest.importorskip("openpyxl")
    import openpyxl
    r = admin_client.get("/delay-report.xlsx")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    detail = wb["Detail"]
    header = [c.value for c in next(detail.iter_rows(min_row=1, max_row=1))]
    so_col = header.index("SO No")
    endpoint_rows = [row for row in detail.iter_rows(min_row=2, values_only=True)
                     if row[so_col] == "NEW-1"]
    assert endpoint_rows, "the downloaded xlsx must also carry rows for NEW-1"


# --- Task 13: the preponed path --- #

@pytest.fixture
def running_optimize(_api_module):
    """A contest already in flight, staged directly rather than run for real (the
    same pattern tests/test_auto_optimize.py and tests/test_manual_apply_backstop.py
    use for `_OPTIMIZE`)."""
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE["state"] = "running"
    yield
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE["state"] = "idle"


@pytest.fixture
def finished_quote_optimize(_api_module, admin_client, uploaded_masters):
    """A completed quote-kind search, staged directly the way
    tests/test_manual_apply_backstop.py stages `_OPTIMIZE` for `_optimize_apply` —
    never a real 10-30 minute contest. Also leaves behind the one draft line
    `/new-orders/prepone` would have, with its typed target date already set (as
    the endpoint would have left it).

    ``quote_movement`` is staged in ``result`` too (2026-09-11 review, I2): the
    real path now computes it ONCE, inside `_finalize_optimize`, and
    `/optimize/status` only ever reads it back from `result` — it is no
    longer recomputed on every poll. A hand-staged `_OPTIMIZE` that skips
    `_finalize_optimize` entirely must stage this the same way a real finished
    contest would have left it, or the status endpoint has nothing to serve."""
    r = admin_client.put("/new-orders/drafts",
                         json={"drafts": [{"so_no": "NEW-1",
                                           "item_code": uploaded_masters, "qty": 25}]})
    assert r.status_code == 200, r.text
    drafts = book_store.load_new_order_drafts()
    for row in drafts:
        row["target_date"] = "2025-03-20"
    book_store.save_new_order_drafts(drafts)
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE.update(
            state="done", kind="quote",
            result={"ranks": {"NEW-1\x1f" + uploaded_masters: 1},
                    "budget": "deep search (new orders)", "seed": 42,
                    "baseline": {"total_late_days": 20, "makespan_days": 6.0},
                    "best": {"total_late_days": 10, "makespan_days": 5.0,
                             "max_committed_slip": 0},
                    "cancelled": False, "improved": True,
                    "best_overlap": None, "current_overlap": None,
                    "flexible_machines": None, "current_flexible": None,
                    "knob": None, "inputs_sig": None,
                    "quote_movement": {"moved": [], "late_days_before": 20,
                                        "late_days_after": 10}})
    return _api_module


def test_prepone_is_admin_only(user_client):
    assert user_client.post("/new-orders/prepone", json={"targets": {}}).status_code == 403


def test_prepone_accept_is_admin_only(user_client):
    assert user_client.post("/new-orders/prepone/accept").status_code == 403


def test_prepone_refuses_when_a_search_is_already_running(admin_client,
                                                          uploaded_masters,
                                                          running_optimize):
    r = admin_client.post("/new-orders/prepone",
                          json={"targets": {"NEW-1\x1f" + uploaded_masters: "2025-03-20"}})
    assert r.status_code == 409


def test_prepone_with_no_drafts_is_a_clear_400(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts", json={"drafts": []})
    r = admin_client.post("/new-orders/prepone", json={"targets": {}})
    assert r.status_code == 400
    assert "no new orders" in r.json()["detail"].lower()


def test_prepone_requires_a_date_for_every_draft_line(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post("/new-orders/prepone", json={"targets": {}})
    assert r.status_code == 400
    assert "NEW-1" in r.json()["detail"]


def test_prepone_refuses_a_bad_date(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post("/new-orders/prepone",
                          json={"targets": {"NEW-1\x1f" + uploaded_masters: "not-a-date"}})
    assert r.status_code == 400


def test_prepone_stores_the_typed_target_on_the_draft(admin_client, uploaded_masters,
                                                      monkeypatch):
    monkeypatch.setattr("api.main._start_optimize", lambda *a, **k: None)
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post("/new-orders/prepone",
                          json={"targets": {"NEW-1\x1f" + uploaded_masters: "2025-03-20"}})
    assert r.status_code == 200
    assert r.json()["started"] is True
    assert book_store.load_new_order_drafts()[0]["target_date"] == "2025-03-20"


def test_prepone_starts_a_quote_kind_search(admin_client, uploaded_masters,
                                            monkeypatch):
    """The search `/new-orders/prepone` starts must be marked `kind="quote"` — the
    Settings panel's ordinary Apply button, and auto-apply, must never touch it.
    Runs the real (fast, 15-eval) search rather than stubbing `_start_optimize`,
    so this also proves `extra_orders` actually reached the contest: the winning
    ranks name the new order's own key, not just the two orders already in the
    book — it competed for a slot like everything else."""
    import api.main as m
    monkeypatch.setitem(m._OPT_BUDGETS, "deep", 15)
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post(
        "/new-orders/prepone",
        json={"targets": {"NEW-1\x1f" + uploaded_masters: "2025-03-20"}})
    assert r.status_code == 200, r.text
    t0 = time.time()
    st = admin_client.get("/optimize/status").json()
    while st["state"] == "running" and time.time() - t0 < 20:
        time.sleep(0.05)
        st = admin_client.get("/optimize/status").json()
    assert st["state"] == "done", st
    assert st["kind"] == "quote"
    with m._OPTIMIZE_LOCK:
        ranks = m._OPTIMIZE["result"]["ranks"]
    assert f"NEW-1\x1f{uploaded_masters}" in ranks


def test_a_finished_quote_searchs_achieved_date_for_the_new_order_is_not_blank(
        admin_client, uploaded_masters, monkeypatch):
    """Review round 2 finding (important): `_metrics_for_ranks` re-read the book
    alone and dropped the in-memory draft line, so the new order's achieved
    date came back blank in `best.expected` every time — the whole point of the
    result screen is comparing what the director typed against what the search
    actually achieved, and without this he would be choosing blind. Runs a
    real (fast) search rather than a hand-staged `_OPTIMIZE`, since the bug
    lived in `_finalize_optimize`'s real recompute path, which a stubbed
    fixture never exercises."""
    import api.main as m
    monkeypatch.setitem(m._OPT_BUDGETS, "deep", 15)
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post(
        "/new-orders/prepone",
        json={"targets": {"NEW-1\x1f" + uploaded_masters: "2025-03-20"}})
    assert r.status_code == 200, r.text
    t0 = time.time()
    st = admin_client.get("/optimize/status").json()
    while st["state"] == "running" and time.time() - t0 < 20:
        time.sleep(0.05)
        st = admin_client.get("/optimize/status").json()
    assert st["state"] == "done", st

    key = f"NEW-1\x1f{uploaded_masters}"
    assert st["best"] is not None
    achieved = st["best"]["expected"].get(key)
    assert achieved, f"achieved date for the new order was blank: {st['best']['expected']}"
    date.fromisoformat(achieved)   # a real, parseable date, not a placeholder

    # The status payload's own before/after summary must agree, and a
    # brand-new order (no "before" to compare against) must never be reported
    # as having "moved" — it was never in the book before this search.
    move = st["quote_movement"]
    assert isinstance(move["late_days_after"], int)
    assert all((row["so_no"], row["item_code"]) != ("NEW-1", uploaded_masters)
              for row in move["moved"])


def test_optimize_status_reports_plan_kind_for_an_ordinary_search(admin_client,
                                                                  uploaded_masters):
    """The ordinary "Start deep search" button must still report kind="plan", and
    must not carry quote_movement — that field is quote-only."""
    import api.main as m
    st = m._start_optimize(budget_evals=15, label="quick", background=False)
    assert st["state"] == "done", st
    assert st["kind"] == "plan"
    assert "quote_movement" not in st or st.get("quote_movement") is None


def test_a_finished_quotes_status_reports_the_movement(admin_client, uploaded_masters,
                                                        finished_quote_optimize):
    """The exact shape the UI task will consume: `kind`, and — once the search is
    `done` — `quote_movement` with `moved` (worst first) and the late-day totals
    before/after, built from `_incumbent_metrics`/`_metrics_for_ranks`, never a
    second comparison."""
    st = admin_client.get("/optimize/status").json()
    assert st["kind"] == "quote"
    assert st["state"] == "done"
    move = st["quote_movement"]
    assert isinstance(move["moved"], list)
    assert isinstance(move["late_days_before"], int)
    assert isinstance(move["late_days_after"], int)


def test_quote_movement_is_computed_once_not_on_every_status_poll(
        admin_client, uploaded_masters, monkeypatch):
    """2026-09-11 review, I2: `_quote_movement` runs two full plan replays
    (`_incumbent_metrics` + `_metrics_for_ranks`). Both roles poll
    `/optimize/status` at boot and on every "Done entering" press, and a
    finished quote result sits in `_OPTIMIZE` until the director accepts or
    discards it — so recomputing on every poll meant seconds of CPU on a
    free-tier instance for a number that cannot change between polls. It must
    now be computed exactly ONCE, inside `_finalize_optimize`, and every poll
    after that must be served straight from the cached `result`."""
    import api.main as m
    calls = []
    real = m._quote_movement

    def _counting(ranks):
        calls.append(1)
        return real(ranks)

    monkeypatch.setattr(m, "_quote_movement", _counting)
    monkeypatch.setitem(m._OPT_BUDGETS, "deep", 15)
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post(
        "/new-orders/prepone",
        json={"targets": {"NEW-1\x1f" + uploaded_masters: "2025-03-20"}})
    assert r.status_code == 200, r.text
    t0 = time.time()
    st = admin_client.get("/optimize/status").json()
    while st["state"] == "running" and time.time() - t0 < 20:
        time.sleep(0.05)
        st = admin_client.get("/optimize/status").json()
    assert st["state"] == "done", st
    assert st["quote_movement"] is not None

    # Poll several more times — a finished quote sits here until accepted or
    # discarded, exactly like a director leaving the tab open, a "Done
    # entering" press elsewhere, or a boot re-fetch after a reload.
    for _ in range(5):
        st2 = admin_client.get("/optimize/status").json()
        assert st2["quote_movement"] == st["quote_movement"]

    assert len(calls) == 1, (
        f"_quote_movement was called {len(calls)} times across the finalize "
        "plus 6 status polls — it must be computed exactly once")


def test_accepting_a_preponed_result_uses_the_typed_date_not_the_achieved_one(
        admin_client, uploaded_masters, finished_quote_optimize):
    r = admin_client.post("/new-orders/prepone/accept")
    assert r.status_code == 200, r.text
    rows = {(o["SO No"], o["Item Code"]): o
            for o in _rows(admin_client.get("/orders").json()["orders"])}
    row = rows[("NEW-1", uploaded_masters)]
    assert row["SO Delivery Date"] == fmt_date(date(2025, 3, 20))


def test_accepting_a_preponed_result_creates_no_queue(admin_client, uploaded_masters,
                                                      finished_quote_optimize):
    """After accept, the queue is empty. This does NOT by itself prove
    accept_prepone's own explicit `book_store.clear_new_order_queue()` line does
    anything: `_optimize_apply()`, which accept always calls, already clears the
    queue on every full optimization. Mutation-tested (2026-09-08 review round
    1): removing accept's own clear line does not fail this or any other test —
    it is honest belt-and-braces, not proof the explicit line is load-bearing."""
    admin_client.post("/new-orders/prepone/accept")
    assert book_store.load_new_order_queue() == []


def test_accepting_a_preponed_result_applies_the_searched_plan(admin_client,
                                                               uploaded_masters,
                                                               finished_quote_optimize):
    """Accepting must adopt the searched sequence (`_optimize_apply()`), not just
    save the orders — the slip figures the director was shown come from that
    sequence, so the applied plan priority must be exactly the ranks the finished
    search staged."""
    admin_client.post("/new-orders/prepone/accept")
    saved = book_store.load_plan_priority()
    assert saved is not None
    assert saved["ranks"] == {"NEW-1\x1f" + uploaded_masters: 1}
    # `_optimize_apply()` also clears the in-memory job so a page refresh can't
    # re-show a stale Apply panel for a plan that is already applied.
    assert admin_client.get("/optimize/status").json()["state"] == "idle"


def test_accepting_clears_the_drafts(admin_client, uploaded_masters,
                                     finished_quote_optimize):
    admin_client.post("/new-orders/prepone/accept")
    assert book_store.load_new_order_drafts() == []


def test_accepting_with_no_drafts_is_a_clear_400(admin_client, uploaded_masters,
                                                 finished_quote_optimize):
    book_store.save_new_order_drafts([])
    r = admin_client.post("/new-orders/prepone/accept")
    assert r.status_code == 400


def test_accepting_refuses_without_a_finished_search(admin_client, uploaded_masters):
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    r = admin_client.post("/new-orders/prepone/accept")
    assert r.status_code == 409


def test_accepting_refuses_an_ordinary_plan_kind_search(admin_client, uploaded_masters,
                                                        _api_module):
    """The Settings panel's ordinary "Start deep search" contest must never be
    adoptable through the Add New Orders accept endpoint — only a `kind="quote"`
    run may be."""
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25,
                                       "target_date": "2025-03-20"}]})
    drafts = book_store.load_new_order_drafts()
    for row in drafts:
        row["target_date"] = "2025-03-20"
    book_store.save_new_order_drafts(drafts)
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE.update(
            state="done", kind="plan",
            result={"ranks": {}, "budget": "deep search", "seed": 42,
                    "baseline": {"total_late_days": 20, "makespan_days": 6.0},
                    "best": {"total_late_days": 10, "makespan_days": 5.0,
                             "max_committed_slip": 0},
                    "cancelled": False, "improved": True,
                    "best_overlap": None, "current_overlap": None,
                    "flexible_machines": None, "current_flexible": None,
                    "knob": None, "inputs_sig": None})
    r = admin_client.post("/new-orders/prepone/accept")
    assert r.status_code == 409


# --- Task 13, review round 1 fixes --- #

def test_accept_refuses_and_changes_nothing_when_the_apply_backstop_rejects(
        admin_client, uploaded_masters):
    """Finding 1 (critical, 2026-09-08 review round 1): `_optimize_apply()`
    carries its own hard guard (the committed-promise backstop) that can reject
    a plan the contest itself scored fine. Apply must run BEFORE any write, so
    a refusal here leaves the book, the drafts and the queue exactly as they
    were — no order added with no plan priority protecting it, no drafts
    wiped, nothing to fix by hand."""
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    drafts = book_store.load_new_order_drafts()
    for row in drafts:
        row["target_date"] = "2025-03-20"
    book_store.save_new_order_drafts(drafts)
    import api.main as m
    with m._OPTIMIZE_LOCK:
        m._OPTIMIZE.update(
            state="done", kind="quote",
            result={"ranks": {"NEW-1\x1f" + uploaded_masters: 1},
                    "budget": "deep search (new orders)", "seed": 42,
                    "baseline": {"total_late_days": 20, "makespan_days": 6.0},
                    # A committed slip far past any slack forces the apply-time
                    # backstop in `_optimize_apply()` to raise 409, AFTER the
                    # search itself already looked like a normal finished quote.
                    "best": {"total_late_days": 10, "makespan_days": 5.0,
                             "max_committed_slip": 999},
                    "cancelled": False, "improved": True,
                    "best_overlap": None, "current_overlap": None,
                    "flexible_machines": None, "current_flexible": None,
                    "knob": None, "inputs_sig": None})
    r = admin_client.post("/new-orders/prepone/accept")
    assert r.status_code == 409
    rows = {(o["SO No"], o["Item Code"])
            for o in _rows(admin_client.get("/orders").json()["orders"])}
    assert ("NEW-1", uploaded_masters) not in rows
    assert book_store.load_new_order_drafts() == drafts


def test_optimize_apply_refuses_a_quote_kind_result(admin_client, uploaded_masters,
                                                     _api_module):
    """Finding 2 (2026-09-08 review round 1): the Settings panel's ordinary
    "Apply this plan" button must never adopt a preponed Add New Orders search
    — that result belongs to its own Accept button."""
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE.update(
            state="done", kind="quote",
            result={"ranks": {"NEW-1\x1f" + uploaded_masters: 1},
                    "budget": "deep search (new orders)", "seed": 42,
                    "baseline": {"total_late_days": 20, "makespan_days": 6.0},
                    "best": {"total_late_days": 10, "makespan_days": 5.0,
                             "max_committed_slip": 0},
                    "cancelled": False, "improved": True,
                    "best_overlap": None, "current_overlap": None,
                    "flexible_machines": None, "current_flexible": None,
                    "knob": None, "inputs_sig": None})
    r = admin_client.post("/optimize/apply")
    assert r.status_code == 409
    assert "Add New Orders" in r.json()["detail"]
    # Unaffected: an ordinary plan-kind result still applies normally.
    with _api_module._OPTIMIZE_LOCK:
        _api_module._OPTIMIZE["kind"] = "plan"
    assert admin_client.post("/optimize/apply").status_code == 200


def test_wiring_quote_movement_passes_the_draft_lines_to_the_after_side(
        admin_client, uploaded_masters, monkeypatch):
    """WIRING CHECK, not a behavioural test: it monkeypatches `_metrics_for_ranks`
    and asserts the `extra_orders` kwarg was passed, so a late-days total
    coincidentally matching either way can't hide a regression. The paired
    BEHAVIOURAL test is
    `test_a_finished_quote_searchs_achieved_date_for_the_new_order_is_not_blank`,
    which checks the actual observable effect (a real achieved date, not a
    call shape). Review round 2 background: `_quote_movement`'s 'after' side
    must be measured on the SAME domain the contest searched (the book plus
    the draft lines), never the saved book alone — that mismatch is what left
    a new order's achieved date blank and could make an EXISTING order's
    recomputed date disagree with what the contest actually found."""
    admin_client.put("/new-orders/drafts",
                     json={"drafts": [{"so_no": "NEW-1",
                                       "item_code": uploaded_masters, "qty": 25}]})
    drafts = book_store.load_new_order_drafts()
    for row in drafts:
        row["target_date"] = "2025-03-20"
    book_store.save_new_order_drafts(drafts)
    import api.main as m
    seen = {}
    real = m._metrics_for_ranks

    def _spy(ranks, *a, **kw):
        seen["extra_orders"] = kw.get("extra_orders")
        return real(ranks, *a, **kw)

    monkeypatch.setattr(m, "_metrics_for_ranks", _spy)
    m._quote_movement({})
    assert seen.get("extra_orders"), (
        "_quote_movement's after-side call did not receive the draft lines")
    assert seen["extra_orders"][0].so_no == "NEW-1"


def test_late_days_totals_are_measured_over_the_existing_book_only(
        admin_client, uploaded_masters, monkeypatch):
    """Review round 3 finding (important): `late_days_before`/`late_days_after`
    must be measured over the SAME set of orders on both sides (the existing
    book) — never each side's own raw total, which sums over whatever domain
    that side happens to cover. "After" covers existing orders plus the new
    draft; "before" covers existing orders only. If the new order is itself
    late against the date the director typed but no existing order moved, the
    two totals must come out EQUAL: the screen's headline number answers "did
    the existing book get pushed", not "does the new order make its date" —
    that second question already has its own per-line answer."""
    import api.main as m
    from engine.models import Order

    key = f"NSO-001\x1f{uploaded_masters}"
    unmoved_iso = "2025-03-15"   # 5 days EARLY against its own 2025-03-20 due date

    monkeypatch.setattr(
        m.book_store, "load_active_orders",
        lambda: {("NSO-001", uploaded_masters): Order(
            so_no="NSO-001", item_code=uploaded_masters, item_name="x",
            ordered_qty=10, delivery_date=date(2025, 3, 20))})
    monkeypatch.setattr(m, "_incumbent_metrics",
                        lambda: {"total_late_days": 0, "expected": {key: unmoved_iso}})

    new_key = f"NEW-1\x1f{uploaded_masters}"

    def _fake_metrics_for_ranks(ranks, *a, **kw):
        # The new order lands badly late against its typed target, and the
        # RAW total (as optimizer.plan_metrics would report it) reflects that
        # — this is exactly the inflated number the fix must not surface.
        return {"total_late_days": 999,
                "expected": {key: unmoved_iso, new_key: "2026-01-01"}}

    monkeypatch.setattr(m, "_metrics_for_ranks", _fake_metrics_for_ranks)
    m.book_store.save_new_order_drafts(
        [{"so_no": "NEW-1", "item_code": uploaded_masters, "qty": 25,
          "target_date": "2025-03-20"}])

    move = m._quote_movement({})
    assert move["moved"] == [], "the unmoved existing order must not appear in moved"
    assert move["late_days_before"] == move["late_days_after"] == 0, move


def test_with_extra_orders_refuses_a_colliding_key():
    """Review round 3 minor item: a draft key colliding with a real order must
    never silently overwrite it. Should be unreachable in practice (the
    drafts endpoint already refuses a duplicate key), but the join itself
    must guard it directly."""
    import api.main as m
    from engine.models import Order
    existing = {("SO1", "X"): Order(so_no="SO1", item_code="X", item_name="real",
                                    ordered_qty=1, delivery_date=date(2025, 1, 1))}
    draft = [Order(so_no="SO1", item_code="X", item_name="draft",
                   ordered_qty=2, delivery_date=date(2025, 2, 1))]
    with pytest.raises(ValueError):
        m._with_extra_orders(existing, draft)
