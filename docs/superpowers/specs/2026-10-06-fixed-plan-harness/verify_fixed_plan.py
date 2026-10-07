"""Fixed-plan verification harness (Task 8, 2026-10-06).

Runs the spec section 5 checks through the REAL API (FastAPI TestClient against
api.main) on a COPY of a store directory. The source directory is never written:
every phase copies it (or a snapshot taken from an earlier phase) into a fresh
scratch directory and points STORE_DIR there. MONGODB_URI / Upstash / the cloud
optimize variables are removed from the environment, so nothing can reach Atlas,
GitHub or a worker.

Credentials: none are stored here. The harness sets throwaway ADMIN_/USER_
USERNAME/PASSWORD env vars before importing api.main (api/auth.py honours them)
and logs in with those.

Simulated days: the app plans from "today" (IST). The harness monkeypatches
``api.main._ist_now`` (``_ist_today`` derives from it) to a controlled clock, so
"day k" is a real calendar date. Each simulated day D (a working day):
  * clock = D 23:00 (so an entry dated D is accepted),
  * every operation the CURRENT plan runs inside [D 08:00, D+1 08:00) is punched
    through POST /actuals: good qty = the planned qty for that day (the whole
    remaining qty when the op ends inside the day, otherwise floor(qty x share of
    its working minutes inside the day)), split over the batch's SO lines in order,
    clipped to the precedence cap (a punch the server would refuse is reduced),
    operator = the plan's operator for that stretch (OS steps without one),
  * on SHORT_DAYS one CNC/VMC op that ends that day gets only 50 % (shortfall) and
    another that continues past the day is punched to its full remaining qty
    (early finish),
  * POST /optimize/done (alternating admin / user role), then POST /run, checks.

Usage (from the repo root, python3.12):
  DEFAULT_SCHEDULER=new python3.12 -B docs/superpowers/specs/2026-10-06-fixed-plan-harness/verify_fixed_plan.py \
      <STORE_DIR_TO_COPY> <OUT_JSON> [--days 10] [--opt-budget 15]
Prints a summary; writes per-check numbers (no order data dump) to OUT_JSON.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import shutil
import sys
import tempfile
import time
from collections import defaultdict
from datetime import date, datetime, timedelta

ap = argparse.ArgumentParser()
ap.add_argument("store")
ap.add_argument("out")
ap.add_argument("--days", type=int, default=10)
ap.add_argument("--opt-budget", type=int, default=15)
ap.add_argument("--start", default="2026-10-06T07:00")
ap.add_argument("--os-on-send", action="store_true",
                help="variant: punch an outsourced step in full on the day its block STARTS "
                     "(parts sent), instead of the day it ends (parts back)")
ARGS = ap.parse_args()

WORK = tempfile.mkdtemp(prefix="fixedplan-verify-")


def fresh_dir(src, name):
    dst = os.path.join(WORK, name)
    shutil.rmtree(dst, ignore_errors=True)
    shutil.copytree(src, dst)
    os.chmod(dst, 0o755)
    for f in os.listdir(dst):
        os.chmod(os.path.join(dst, f), 0o644)
    return dst


for k in ("MONGODB_URI", "UPSTASH_REDIS_REST_URL", "UPSTASH_REDIS_REST_TOKEN",
          "GITHUB_DISPATCH_TOKEN", "OPTIMIZE_WORKER_SECRET"):
    os.environ.pop(k, None)
_ADMIN = ("verify-admin", "verify-admin-" + os.urandom(6).hex())
_USER = ("verify-user", "verify-user-" + os.urandom(6).hex())
os.environ.update(DEFAULT_SCHEDULER="new", AUTO_OPTIMIZE="0",
                  ADMIN_USERNAME=_ADMIN[0], ADMIN_PASSWORD=_ADMIN[1],
                  USER_USERNAME=_USER[0], USER_PASSWORD=_USER[1],
                  STORE_DIR=fresh_dir(ARGS.store, "phase-main"))
sys.path.insert(0, os.getcwd())

from fastapi.testclient import TestClient  # noqa: E402
from api import main as m  # noqa: E402
from engine import book_store, optimizer, orderbook  # noqa: E402
KEY_SEP = "\x1f"

CLOCK = {"now": datetime.fromisoformat(ARGS.start)}
m._ist_now = lambda: CLOCK["now"]
OS_LANES = {"OS / Outsourced", "Off-machine"}
RESULTS: dict = {"checks": [], "days": []}


def check(name, ok, **nums):
    RESULTS["checks"].append({"name": name, "pass": bool(ok), **nums})
    print(("PASS " if ok else "FAIL ") + name + "  " + json.dumps(nums, default=str)[:400])
    return ok


def use_store(path):
    os.environ["STORE_DIR"] = path
    m._PLAN_CACHE["key"] = None
    m._PLAN_CACHE["result"] = None
    with m._OPTIMIZE_LOCK:
        m._OPTIMIZE.update(state="idle", result=None, best=None, baseline=None)


def snapshot(name):
    return fresh_dir(os.environ["STORE_DIR"], name)


def clients():
    a, u = TestClient(m.app), TestClient(m.app)
    assert a.post("/login", data={"username": _ADMIN[0], "password": _ADMIN[1]},
                  follow_redirects=False).status_code == 303
    assert u.post("/login", data={"username": _USER[0], "password": _USER[1]},
                  follow_redirects=False).status_code == 303
    return a, u


ADMIN, USER = clients()


# ----------------------------------------------------------------- plan access
def run(client=None):
    r = (client or ADMIN).post("/run", json={})
    assert r.status_code == 200, r.text[:300]
    j = r.json()
    art = m._PLAN_CACHE["artifacts"]
    return j, art["plan_run"], art["so_lines"], art["config"], art["masters"]


def plan_hash(schedule):
    rows = sorted((e.machine, e.operator or "", e.start.isoformat(), e.end.isoformat(),
                   float(e.qty), e.item_code, e.process_seq, tuple(sorted(e.so_refs or [])),
                   tuple((s.isoformat(), en.isoformat(), n or "") for s, en, n in (e.op_segments or [])))
                  for e in schedule)
    return hashlib.sha256(json.dumps(rows, default=str).encode()).hexdigest()[:16]


def late_days(schedule, so_lines, config):
    pm = optimizer.plan_metrics(schedule, so_lines, config.plan_start_date,
                                promise_slack_days=getattr(config, "committed_promise_slack_days", 3))
    return round(float(pm["total_late_days"]), 1), round(float(pm["makespan_days"]), 2)


def step_machines(schedule):
    """(so, item, seq) -> set of machines, real machine ops only."""
    out = defaultdict(set)
    for e in schedule:
        if e.machine in OS_LANES:
            continue
        for so in (e.piece_refs or e.so_refs or []):
            out[(so, e.item_code, e.process_seq)].add(e.machine)
    return out


def published_turns(rows):
    """Per machine, the published rows in START order (the turn, spec 4.1)."""
    by_m = defaultdict(list)
    for i, r in enumerate(rows):
        by_m[r["machine"]].append(i)
    turn = {}
    for mid, idx in by_m.items():
        idx.sort(key=lambda i: (rows[i]["start"], rows[i].get("placed", 0)))
        for t, i in enumerate(idx):
            turn[i] = t
    return turn


def machine_and_turn_audit(schedule, rows, frozen_keys):
    """0 ops on another machine than published; per machine, the remaining ops against
    the published order. Since the owner's amendment (spec section 8, Task 9) a later
    job may run before an earlier one when the earlier one cannot start yet, so every
    inversion (later-turn job B starts before earlier-turn job A on one machine) is
    classified:
      * frozen      : either job is half finished (frozen ops run first),
      * not_ready   : A's previous step (in its batch, OS included) was still running
                      when B started, i.e. A could not start then (next ready job),
      * ready_waited: A's previous step had ended by B's start AND A starts the moment
                      B ends, i.e. A was ready and B took the slot in front of it.
                      This is what the rule forbids ("when several are ready, the
                      published order wins"); must be 0,
      * placed_first: the published plan placed B before A (it laid A into an
                      earlier gap afterwards); the repair replays placement order
                      (Task 6b) and the gap no longer holds A. Allowed; reported,
      * ready_hole  : A was ready but did not start when B ended: B used time on the
                      machine that A could not use (no person, or a gap A's run did
                      not fit). Allowed; reported.
    The not-ready test ignores overlap (A may start before its previous step ends), so
    it is lenient towards the engine; the engine tests pin the exact rule."""
    turn = published_turns(rows)
    index = defaultdict(list)          # (item, seq, so) -> [row idx]
    for i, r in enumerate(rows):
        for so in r["so_refs"]:
            index[(r["item_code"], r["process_seq"], so)].append(i)
    steps = defaultdict(lambda: defaultdict(list))   # (batch, item) -> seq -> entries
    for e in schedule:
        steps[(e.batch_id, e.item_code)][e.process_seq].append(e)

    def prev_end(e):
        # The paced end of everything before A in its routing: a step never finishes
        # before the steps ahead of it, so the latest of them gates A (not only the
        # immediately previous step, which can end earlier than one before it).
        ends = [x.end for q, xs in steps[(e.batch_id, e.item_code)].items()
                if q < e.process_seq for x in xs]
        return max(ends) if ends else None

    moved, unpublished = [], 0
    cls = {"frozen": [], "not_ready": [], "placed_first": [], "ready_waited": [],
           "ready_hole": []}
    seq_by_m = defaultdict(list)
    # A step laid AROUND other jobs on a manual station is several entries (and
    # several published rows): the k-th entry of a step on a machine takes the turn
    # of the k-th published row of that step there, not the first row's turn (that
    # read the second part as jumping every job laid in between).
    part = defaultdict(int)
    for e in sorted(schedule, key=lambda x: x.start):
        if e.machine in OS_LANES:
            continue
        refs = e.piece_refs or e.so_refs or []
        cands = {i for so in refs for i in index.get((e.item_code, e.process_seq, so), [])}
        if not cands:
            unpublished += 1
            continue
        same = [i for i in cands if rows[i]["machine"] == e.machine]
        if not same:
            moved.append((refs[:2], e.item_code, e.process_seq, e.machine,
                          sorted({rows[i]["machine"] for i in cands})))
            continue
        same.sort(key=lambda i: (rows[i]["start"], rows[i].get("placed", 0)))
        pk = (e.machine, e.item_code, e.process_seq, tuple(sorted(refs)))
        k = part[pk]
        part[pk] += 1
        row = same[min(k, len(same) - 1)]
        seq_by_m[e.machine].append((e.start, turn[row], e, rows[row].get("placed")))
    for mid, lst in seq_by_m.items():
        lst.sort(key=lambda x: (x[0], x[1]))
        best = -1
        prev = prev_placed = None
        for st, t, e, placed in lst:
            if t < best:
                ex = (mid, e.item_code, e.process_seq, t, best)
                fz = any((so, e.item_code, e.process_seq) in frozen_keys
                         for so in (e.piece_refs or e.so_refs or [])) or \
                     any((so, prev.item_code, prev.process_seq) in frozen_keys
                         for so in (prev.piece_refs or prev.so_refs or []))
                pe = prev_end(e)
                # B's step may be several entries (a manual job running past the end
                # of the shift into the next day): what matters is when B STARTED,
                # i.e. its first entry, not the part that resumed in front of A.
                b_start = min(x.start for x in steps[(prev.batch_id, prev.item_code)]
                              [prev.process_seq] if x.machine == prev.machine)
                if fz:
                    cls["frozen"].append(ex)
                elif pe is not None and pe > b_start:
                    cls["not_ready"].append(ex)
                elif (isinstance(placed, int) and isinstance(prev_placed, int)
                      and prev_placed < placed):
                    # The published plan PLACED B first and laid A into an earlier
                    # gap; the repair replays the placement order (Task 6b), and
                    # that gap no longer holds A. Published order, by placement.
                    cls["placed_first"].append(ex)
                elif e.start == prev.end:
                    cls["ready_waited"].append(ex + (
                        "A", str(e.start), "A prev step end", str(pe), "A placed", placed,
                        "B", prev.item_code, prev.process_seq, str(prev.start), str(prev.end),
                        "B placed", prev_placed))
                else:
                    cls["ready_hole"].append(ex)
            if t >= best:
                best, prev, prev_placed = t, e, placed
    return {"machine_changed": len(moved), "machine_changed_ex": moved[:5],
            "turn_inversions": len(cls["ready_waited"]),
            "turn_inversions_ex": cls["ready_waited"][:5],
            "turn_inversions_frozen": len(cls["frozen"]),
            "turn_inversions_frozen_ex": cls["frozen"][:5],
            "inversions_not_ready": len(cls["not_ready"]),
            "inversions_ready_hole": len(cls["ready_hole"]),
            "inversions_placed_first": len(cls["placed_first"]),
            "inversions_ready_hole_ex": cls["ready_hole"][:5],
            "unpublished_ops": unpublished}


def held_audit(schedule, rows, frozen_keys, tol_min=1):
    """Fix round 2 (re-review): an ON-TIME job A whose machine was HELD by a later-
    published job B at the moment A became ready is a violation, unless B was the
    give-way of that machine for a LATE job L (published before B, not ready when B
    started). "Ready" is the paced end of A's earlier steps (overlap ignored, so it
    can only be later than the engine's); "on time" is ready <= A's published start
    + 1 minute; published order is the rows' placement order. Steps are read by their
    first entry on the machine; frozen pairs are left out."""
    ps = min((e.start for e in schedule), default=None)
    steps = defaultdict(lambda: defaultdict(list))
    for e in schedule:
        steps[(e.batch_id, e.item_code)][e.process_seq].append(e)
    index = defaultdict(list)
    for i, r in enumerate(rows):
        for so in r["so_refs"]:
            index[(r["item_code"], r["process_seq"], so)].append(i)
    tol = timedelta(minutes=tol_min)
    per_m = defaultdict(list)          # machine -> [(first start, end, ready, pub, placed, entry)]
    seen = set()
    for e in schedule:
        if e.machine in OS_LANES:
            continue
        sk = (e.batch_id, e.item_code, e.process_seq, e.machine)
        if sk in seen:
            continue
        seen.add(sk)
        parts = [x for x in steps[(e.batch_id, e.item_code)][e.process_seq] if x.machine == e.machine]
        refs = e.piece_refs or e.so_refs or []
        cand = [i for so in refs for i in index.get((e.item_code, e.process_seq, so), [])
                if rows[i]["machine"] == e.machine]
        if not cand:
            continue
        row = min(cand, key=lambda i: (rows[i]["start"], rows[i].get("placed", 0)))
        earlier = [x.end for q, xs in steps[(e.batch_id, e.item_code)].items()
                   if q < e.process_seq for x in xs]
        ready = max(earlier) if earlier else ps
        fz = any((so, e.item_code, e.process_seq) in frozen_keys for so in refs)
        per_m[e.machine].append((min(x.start for x in parts), max(x.end for x in parts), ready,
                                 datetime.fromisoformat(rows[row]["start"]),
                                 rows[row].get("placed"), e, fz))
    bad = []
    for mid, lst in per_m.items():
        for a_st, a_en, a_rd, a_pub, a_pl, a, a_fz in lst:
            if a_fz or a_pl is None or a_rd > a_pub + tol or a_st <= a_rd:
                continue
            for b_st, b_en, b_rd, b_pub, b_pl, b, b_fz in lst:
                if b is a or b_fz or b_pl is None or b_pl <= a_pl:
                    continue
                if not (b_st <= a_rd < b_en and a_st >= b_en - tol):
                    continue
                give_way = any(l_pl is not None and l_pl < b_pl and l_rd > l_pub + tol
                               and l_rd > b_st and not l_fz
                               for _s, _e, l_rd, l_pub, l_pl, l, l_fz in lst if l is not b)
                if not give_way:
                    bad.append((mid, a.item_code, a.process_seq, str(a_rd), b.item_code,
                                b.process_seq, str(b_st), str(b_en)))
    return bad


# ------------------------------------------------------------ invariant checks
def shift_day(t: datetime) -> date:
    return (t - timedelta(days=1)).date() if t.hour < 5 else t.date()


def calendar_people_audit(schedule, masters, config, downtime, absences):
    cal = masters.calendar
    thr = getattr(config, "two_shift_threshold_hours", 12.0)
    two = {mid: mc.is_two_shift(thr) for mid, mc in masters.machines.items()}
    dt_days = defaultdict(set)
    for d in downtime:
        a, b = date.fromisoformat(d["from_date"]), date.fromisoformat(d["to_date"])
        while a <= b:
            dt_days[d["machine"]].add(a)
            a += timedelta(days=1)
    absent = defaultdict(list)
    for a in absences:
        absent[a["operator"]].append((datetime.combine(date.fromisoformat(a["from_date"]), datetime.min.time()),
                                      datetime.combine(date.fromisoformat(a["to_date"]) + timedelta(days=1),
                                                       datetime.min.time())))
    bad = defaultdict(list)
    book = defaultdict(list)
    for e in schedule:
        if e.machine in OS_LANES:
            continue
        for s, en, who in (e.op_segments or []):
            if en <= s:
                continue
            sd = shift_day(s)
            # a segment never crosses a shift boundary or a meal break
            for day in (s.date() - timedelta(days=1), s.date()):
                for bs, be in ((13 * 60, 13 * 60 + 30), (22 * 60, 22 * 60 + 30)):
                    b0 = datetime.combine(day, datetime.min.time()) + timedelta(minutes=bs)
                    b1 = datetime.combine(day, datetime.min.time()) + timedelta(minutes=be)
                    if s < b1 and en > b0:
                        bad["meal_break"].append((e.machine, s.isoformat()))
            if not cal.is_working_day(sd):
                bad["weekly_off_or_holiday"].append((e.machine, s.isoformat()))
            if sd in dt_days.get(e.machine, ()):
                bad["down_machine"].append((e.machine, s.isoformat()))
            first0 = datetime.combine(sd, datetime.min.time()) + timedelta(hours=8)
            first1 = datetime.combine(sd, datetime.min.time()) + timedelta(hours=19)
            sec1 = datetime.combine(sd, datetime.min.time()) + timedelta(hours=29)
            if s < first0 or en > sec1:
                bad["outside_shifts"].append((e.machine, s.isoformat()))
            if not two.get(e.machine, True) and en > first1:
                bad["night_on_single_shift"].append((e.machine, s.isoformat()))
            if who:
                book[who].append((s, en, e.machine, e.batch_id, e.process_seq))
                for a0, a1 in absent.get(who, ()):
                    if s < a1 and en > a0:
                        bad["absent_booked"].append((who, s.isoformat()))
    dbl = 0
    ex = []
    for who, lst in book.items():
        lst.sort()
        for (s1, e1, m1, b1, p1), (s2, e2, m2, b2, p2) in zip(lst, lst[1:]):
            if s2 < e1 and (b1, p1, m1) != (b2, p2, m2):
                dbl += 1
                ex.append((who, m1, m2, s2.isoformat()))
    out = {k: len(v) for k, v in bad.items()}
    out["double_booked"] = dbl
    out["examples"] = {k: v[:3] for k, v in bad.items()}
    out["double_ex"] = ex[:3]
    return out


def report_kinds(j):
    rep = j["report"]
    cols = rep["columns"]
    if "Kind" not in cols:
        return {}
    k = cols.index("Kind")
    cnt = defaultdict(int)
    for r in rep["rows"]:
        cnt[r[k]] += 1
    return dict(cnt)


def surfaces_audit(j):
    """Orders (expected_end) vs Gantt vs shift-wise vs delay report."""
    exp = {tuple(k.split(KEY_SEP)): v for k, v in j["expected_end"].items()}
    mism = defaultdict(int)
    ex = []
    g = ADMIN.get("/gantt").json()
    for row in g["rows"]:
        d = datetime.strptime(row["completion"], "%d-%m-%Y").date().isoformat()
        for so in [s.strip() for s in row["so_no"].split(",") if s.strip()]:
            key = (so, row["item_code"])
            if key in exp and exp[key] != d:
                mism["gantt"] += 1
                ex.append(("gantt", key, exp[key], d))
    sw = j["trace"]["rule6"].get("shiftwise") or {"columns": [], "rows": []}
    if sw["rows"]:
        c = sw["columns"]
        iso, iitem, iexp = c.index("SO No"), c.index("Item Code"), c.index("Expected completion")
        for r in sw["rows"]:
            d = datetime.strptime(r[iexp], "%d-%m-%Y").date().isoformat()
            for so in [s.strip() for s in r[iso].split(",") if s.strip()]:
                key = (so, r[iitem])
                if key in exp and exp[key] != d:
                    mism["shiftwise"] += 1
                    ex.append(("shiftwise", key, exp[key], d))
    from openpyxl import load_workbook
    resp = USER.get("/delay-report.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    iso, iitem, iexp = hdr.index("SO No"), hdr.index("Item Code"), hdr.index("Expected Completion")
    seen = set()
    for r in rows[1:]:
        key = (r[iso], r[iitem])
        seen.add(key)
        d = datetime.strptime(r[iexp], "%d-%m-%Y").date().isoformat()
        if key in exp and exp[key] != d:
            mism["delay_report"] += 1
            ex.append(("delay", key, exp[key], d))
    mism["delay_missing"] = len(set(exp) - seen)
    return {"orders": len(exp), **dict(mism), "examples": ex[:4]}


def full_audit(tag, j, plan_run, so_lines, config, masters, rows, frozen_keys):
    sched = plan_run.schedule
    a = machine_and_turn_audit(sched, rows, frozen_keys)
    c = calendar_people_audit(sched, masters, config, book_store.load_machine_downtime(),
                              book_store.load_absences())
    k = report_kinds(j)
    s = surfaces_audit(j)
    inv = {x: k.get(x, 0) for x in ("ROUTING_ORDER_VIOLATION", "OPERATOR_NOT_QUALIFIED",
                                    "BATCH_QTY_SHORT")}
    ok = (a["machine_changed"] == 0 and a["turn_inversions"] == 0)
    check(f"{tag}: machines held, published order wins among ready jobs", ok,
          **{x: a[x] for x in ("machine_changed", "turn_inversions",
                               "turn_inversions_frozen", "inversions_not_ready",
                               "inversions_ready_hole", "inversions_placed_first",
                               "unpublished_ops")},
          ex=a["machine_changed_ex"][:2] + a["turn_inversions_ex"][:2])
    cal_keys = ("meal_break", "weekly_off_or_holiday", "down_machine", "outside_shifts",
                "night_on_single_shift")
    held = held_audit(sched, rows, frozen_keys)
    check(f"{tag}: no on-time ready job finds its machine held by a later job", not held,
          held=len(held), ex=held[:3])
    check(f"{tag}: calendar", all(c.get(x, 0) == 0 for x in cal_keys),
          **{x: c.get(x, 0) for x in cal_keys}, ex=c["examples"])
    check(f"{tag}: people", c["double_booked"] == 0 and c.get("absent_booked", 0) == 0,
          double_booked=c["double_booked"], absent_booked=c.get("absent_booked", 0),
          ex=c["double_ex"])
    check(f"{tag}: invariants", sum(inv.values()) == 0, **inv)
    check(f"{tag}: one set of dates", all(s.get(x, 0) == 0 for x in
                                         ("gantt", "shiftwise", "delay_report", "delay_missing")),
          **{x: s.get(x, 0) for x in ("orders", "gantt", "shiftwise", "delay_report",
                                      "delay_missing")}, ex=s["examples"])
    return a, c, inv, s


# ------------------------------------------------------------- punch simulator
def work_minutes(segs, lo=None, hi=None):
    t = 0.0
    for s, en, _w in segs:
        a = max(s, lo) if lo else s
        b = min(en, hi) if hi else en
        if b > a:
            t += (b - a).total_seconds() / 60.0
    return t


def first_operator(e, lo, hi):
    for s, en, w in (e.op_segments or []):
        if w and en > lo and s < hi:
            return w
    return e.operator or ""


def day_plan(schedule, D):
    lo = datetime.combine(D, datetime.min.time()) + timedelta(hours=8)
    hi = lo + timedelta(days=1)
    out = []
    for e in schedule:
        qty = float(e.qty or 0)
        if qty <= 0:
            continue
        segs = e.op_segments or []
        if e.machine in OS_LANES or not segs:
            # OS: a flat continuous block whose length does not depend on the qty, so
            # the parts are punched when they come back (the block ends inside the
            # day), never prorated: a prorated OS punch would make the engine lay a
            # fresh full-length block from the next morning. Off-machine milestone:
            # zero length, punched on its day.
            when = e.start if (ARGS.os_on_send and e.machine == "OS / Outsourced") else e.end
            if lo <= when < hi or (e.end == e.start and lo <= e.start < hi):
                out.append((e, qty, True, ""))
            continue
        tot = work_minutes(segs)
        inday = work_minutes(segs, lo, hi)
        if inday <= 0:
            continue
        ends_today = max(en for _s, en, _w in segs) <= hi
        q = qty if ends_today else math.floor(qty * inday / tot)
        if q > 0:
            out.append((e, q, ends_today, first_operator(e, lo, hi)))
    out.sort(key=lambda x: (x[0].start, x[0].process_seq))
    return out


def room(e, masters):
    """How many pieces of op ``e`` the server would accept right now (precedence cap)."""
    active = book_store.load_active_orders()
    actuals = book_store.load_actuals()
    routing = masters.routings.get(e.item_code)
    names = [orderbook._norm(p.name) for p in sorted(routing.processes, key=lambda p: p.seq)] if routing else []
    tgt = orderbook._norm(e.process_name)
    tot = 0.0
    for so in (e.piece_refs or e.so_refs or []):
        o = active.get((so, e.item_code))
        if o is None:
            continue
        produced, good = orderbook._process_totals(actuals, so, e.item_code)
        cap = o.ordered_qty if (not names or tgt not in names or names.index(tgt) == 0) \
            else good.get(names[names.index(tgt) - 1], 0.0)
        tot += max(0.0, min(o.ordered_qty, cap) - produced.get(tgt, 0.0))
    return tot


def punch(e, q, D, operator, masters, client):
    """Split one op's punch over its SO lines, clipped to what the server accepts."""
    active = book_store.load_active_orders()
    actuals = book_store.load_actuals()
    routing = masters.routings.get(e.item_code)
    names = [orderbook._norm(p.name) for p in sorted(routing.processes, key=lambda p: p.seq)] if routing else []
    tgt = orderbook._norm(e.process_name)
    left, sent, clipped = q, 0, 0
    refs = sorted(e.piece_refs or e.so_refs or [],
                  key=lambda so: (active[(so, e.item_code)].delivery_date if (so, e.item_code) in active else date.max, so))
    if not operator and not orderbook.process_is_outsourced(routing, e.process_name):
        operator = (masters.operators[0].name if masters.operators else "")
    for so in refs:
        if left <= 0:
            break
        o = active.get((so, e.item_code))
        if o is None:
            continue
        produced, good = orderbook._process_totals(actuals, so, e.item_code)
        got = produced.get(tgt, 0.0)
        cap = o.ordered_qty if (not names or tgt not in names or names.index(tgt) == 0) \
            else good.get(names[names.index(tgt) - 1], 0.0)
        room = max(0.0, min(o.ordered_qty, cap) - got)
        n = math.floor(min(left, room))
        if n <= 0:
            continue
        r = client.post("/actuals", json={
            "so_no": so, "item_code": e.item_code, "entry_date": D.isoformat(),
            "shift": "1st shift", "process": e.process_name, "operator": operator,
            "qty_produced": n, "machine": "" if e.machine in OS_LANES else e.machine})
        if r.status_code != 200:
            continue
        actuals = book_store.load_actuals()
        left -= n
        sent += n
    return sent, max(0, math.floor(q) - sent)


def downstream_orders(schedule, sources, people=False):
    """Orders reachable from the ``sources`` ops: later on the same machine, later
    steps of the same batch (and, with ``people``, later work of the same person),
    transitively. Spec 5.5's 'shares a machine queue'."""
    by_m, by_b, by_p = defaultdict(list), defaultdict(list), defaultdict(list)
    for e in schedule:
        by_m[e.machine].append(e)
        by_b[e.batch_id].append(e)
        for _s, _e, w in (e.op_segments or []):
            if w:
                by_p[w].append(e)
    seen, stack = set(), list(sources)
    while stack:
        x = stack.pop()
        if id(x) in seen:
            continue
        seen.add(id(x))
        if x.machine not in OS_LANES:
            stack += [y for y in by_m[x.machine] if y.start >= x.start]
        stack += [y for y in by_b[x.batch_id] if y.process_seq >= x.process_seq]
        if people:
            for _s, _e, w in (x.op_segments or []):
                if w:
                    stack += [y for y in by_p[w] if y.start >= x.start]
    return {(so, e.item_code) for e in schedule if id(e) in seen for so in (e.so_refs or [])}


def dates_of(j):
    return dict(j["expected_end"])


def moved(a, b):
    return {k: (a[k], b[k]) for k in a if k in b and a[k] != b[k]}


def frozen_keys_now():
    return {(f["so_no"], f["item_code"], f["op_seq"]) for f in book_store.load_frozen_ops()}


def next_working(d, cal):
    d += timedelta(days=1)
    while not cal.is_working_day(d):
        d += timedelta(days=1)
    return d


# =================================================================== PHASE 1
t0 = time.time()
cfg0 = m._load_plan_config()
# (a) BEFORE this feature: the free plan with the applied ranks (what main serves).
m._PLAN_CACHE["key"] = None
m._plan(cfg0, _seeding=True)
art = m._PLAN_CACHE["artifacts"]
free0 = art["plan_run"].schedule
free0_dates = {f"{so}{KEY_SEP}{it}": d.isoformat()
               for (so, it), d in optimizer.expected_completion(free0).items()}
late_a = late_days(free0, art["so_lines"], art["config"])
m._PLAN_CACHE["key"] = None
assert not book_store.load_published_meta(), "the copy already has a published plan"

# Go-live: the first /run seeds the published plan, then repairs it.
j0, pr0, sl0, c0, ms0 = run()
rows0 = book_store.load_last_applied_schedule()
late_b = late_days(pr0.schedule, sl0, c0)
d0 = dates_of(j0)
mv = moved(free0_dates, d0)
RESULTS["golive"] = {"published_rows": len(rows0), "orders": len(d0),
                     "late_free": late_a, "late_repaired": late_b,
                     "orders_moved": len(mv),
                     "max_move_days": max((abs((date.fromisoformat(b) - date.fromisoformat(a)).days)
                                           for a, b in mv.values()), default=0),
                     "plan_start": j0["resolved_plan_start"], "alerts": j0["fixed_plan_alerts"]}
print("GO-LIVE", RESULTS["golive"])
full_audit("day 0 (go-live)", j0, pr0, sl0, c0, ms0, rows0, frozen_keys_now())

# Idempotence (spec 5.4): Done twice, no punch in between.
CLOCK["now"] = datetime.fromisoformat(ARGS.start)
r1 = ADMIN.post("/optimize/done")
ja, pa, *_ = run()
r2 = USER.post("/optimize/done")
jb, pb, *_ = run()
check("idempotence day 0: Done twice, identical plan",
      r1.status_code == r2.status_code == 200 and plan_hash(pa.schedule) == plan_hash(pb.schedule),
      first_run=plan_hash(pr0.schedule), done1=plan_hash(pa.schedule), done2=plan_hash(pb.schedule),
      reason=r1.json().get("reason"), dates_moved_by_first_done=len(moved(d0, dates_of(ja))))
check("Done never starts a search", m._OPTIMIZE["state"] == "idle"
      and r1.json().get("started") is False, state=m._OPTIMIZE["state"])
note = (book_store.load_auto_note() or {}).get("text", "")
check("Done note names who pressed and says only times changed",
      _USER[0] in note and "only times changed" in note, note=note[:160])
SNAP_DAY0 = snapshot("snap-day0")

# =================================================================== PHASE 2
SHORT_DAYS = {3, 7}
cal = ms0.calendar
D = date.fromisoformat(jb["resolved_plan_start"])
cur_j, cur_pr = jb, pb
for k in range(1, ARGS.days + 1):
    CLOCK["now"] = datetime.combine(D, datetime.min.time()) + timedelta(hours=23)
    sched = cur_pr.schedule
    before = dates_of(cur_j)
    plan = day_plan(sched, D)
    short_e = early_e = None
    if k in SHORT_DAYS:
        cnc = [x for x in plan if x[0].machine[:3] in ("CNC", "VMC")]
        ending = [x for x in cnc if x[2] and x[1] >= 4]
        going = [x for x in cnc if not x[2] and room(x[0], ms0) >= float(x[0].qty)]
        if ending:
            short_e = ending[len(ending) // 2][0]
        if going:
            early_e = going[0][0]
    sent = clipped = 0
    deviating, rounding = [], []
    for e, q, _ends, who in plan:
        intended = q
        if e is short_e:
            intended = math.floor(q / 2)
        if e is early_e:
            intended = float(e.qty)
        s_, c = punch(e, intended, D, who, ms0, ADMIN)
        sent += s_
        clipped += c
        if s_ != q:
            deviating.append(e)          # short, early, or the server's cap clipped it
        elif not _ends and e.machine not in OS_LANES:
            rounding.append(e)           # whole pieces only: the plan's day share was fractional
    # OS blocks still at the vendor are re-laid in full from the next plan start
    # (flat duration, pre-existing engine behaviour): a source of movement too.
    hi = datetime.combine(D, datetime.min.time()) + timedelta(hours=32)
    os_open = [] if ARGS.os_on_send else [
        e for e in sched if e.machine == "OS / Outsourced" and e.end >= hi and e.start < hi]
    by = USER if k % 2 else ADMIN
    rd = by.post("/optimize/done")
    j, pr, sl, cf, ms = run(by)
    after = dates_of(j)
    mvd = moved(before, after)
    moved_keys = {tuple(x.split(KEY_SEP)) for x in mvd}
    # An order with nothing left to make at any step, but not finished at its gate
    # (so not archived), is a row of zero-length milestones at the plan start: its
    # date follows the plan clock every day. Not a placement effect; set aside.
    live_qty = defaultdict(float)
    for e in pr.schedule:
        for so in (e.so_refs or []):
            live_qty[(so, e.item_code)] += float(e.qty or 0)
    nothing_left = {kk for kk in moved_keys if live_qty.get(kk, 0.0) == 0.0}
    moved_keys -= nothing_left
    expl_dev = downstream_orders(sched, deviating)
    expl_all = downstream_orders(sched, deviating + os_open + rounding)
    expl_ppl = downstream_orders(sched, deviating + os_open + rounding, people=True)
    day = {"day": k, "date": D.isoformat(), "ops_punched": len(plan), "pieces": sent,
           "pieces_short_of_plan": clipped, "done": rd.status_code,
           "plan_start_after": m._PLAN_CACHE["artifacts"]["config"].plan_start_date.isoformat(),
           "orders_moved": len(mvd),
           "max_move": max((abs((date.fromisoformat(b) - date.fromisoformat(a)).days)
                            for a, b in mvd.values()), default=0),
           "late": late_days(pr.schedule, sl, cf), "frozen": len(book_store.load_frozen_ops()),
           "alerts": j["fixed_plan_alerts"],
           "moved_nothing_left_to_make": len(nothing_left),
           "deviating_ops": len(deviating), "os_blocks_open": len(os_open),
           "rounding_ops": len(rounding),
           "moved_outside_downstream_of_deviations": len(moved_keys - expl_dev),
           "moved_outside_downstream_incl_os_and_rounding": len(moved_keys - expl_all),
           "moved_outside_incl_people": len(moved_keys - expl_ppl),
           "outside_ex": sorted(moved_keys - expl_ppl)[:4]}
    if short_e is not None:
        day.update(short=(short_e.machine, short_e.item_code, short_e.process_seq),
                   early=(early_e.machine, early_e.item_code, early_e.process_seq) if early_e else None)
    RESULTS["days"].append(day)
    print("DAY", json.dumps(day, default=str)[:600])
    full_audit(f"day {k} ({D})", j, pr, sl, cf, ms, rows0, frozen_keys_now())
    if k == 5:
        h1 = plan_hash(pr.schedule)
        ADMIN.post("/optimize/done")
        jx, px, *_ = run()
        check("idempotence day 5: a second Done with no punch", plan_hash(px.schedule) == h1,
              a=h1, b=plan_hash(px.schedule))
        j, pr = jx, px
    cur_j, cur_pr = j, pr
    D = next_working(D, cal)

DAY10_PLAN_START = cur_j["resolved_plan_start"]
late_c = late_days(cur_pr.schedule, m._PLAN_CACHE["artifacts"]["so_lines"],
                   m._PLAN_CACHE["artifacts"]["config"])
m._PLAN_CACHE["key"] = None
m._plan(m._load_plan_config(), _seeding=True)       # the free plan on day 10, same ranks
art = m._PLAN_CACHE["artifacts"]
late_c_free = late_days(art["plan_run"].schedule, art["so_lines"], art["config"])
m._PLAN_CACHE["key"] = None
SNAP_DAY10 = snapshot("snap-day10")
OPT_CLOCK = CLOCK["now"]

# =================================================================== PHASE 3
# (d) a fresh Optimize on day 10 (local, small budget), Apply, and the list's truth.
use_store(fresh_dir(SNAP_DAY10, "phase-opt"))
run()
t = time.time()
m._start_optimize(ARGS.opt_budget, "deep", background=False)
st = ADMIN.get("/optimize/status").json()
opt_secs = round(time.time() - t, 1)
best = st.get("best") or (st.get("result") or {}).get("best") or {}
dc = st.get("date_changes")
before = dates_of(run()[0])
ra = ADMIN.post("/optimize/apply")
ja, pa, sla, cfa, msa = run()
after = dates_of(ja)
listed = {f"{r['so']}{KEY_SEP}{r['item']}": r["after"] for r in (dc or [])}
wrong = [k for k, v in listed.items() if after.get(k) != v]
unlisted_moved = [k for k in moved(before, after) if k not in listed]
late_d = late_days(pa.schedule, sla, cfa)
inc_after = m._incumbent_metrics()["total_late_days"]
check("Optimize day 10: the panel's After late-days is the plan after Apply (final review I2)",
      best.get("total_late_days") == inc_after,
      panel=best.get("total_late_days"), after_apply=inc_after)
RESULTS["optimize_day10"] = {"state": st.get("state"), "secs": opt_secs,
                             "search_best_late": best.get("total_late_days"),
                             "date_changes": None if dc is None else len(dc),
                             "apply": ra.status_code, "late_after_apply": late_d}
check("Optimize day 10: listed dates are exactly what the floor gets after Apply",
      ra.status_code == 200 and dc is not None and not wrong and not unlisted_moved,
      listed=len(listed), wrong=len(wrong), unlisted_moved=len(unlisted_moved),
      ex=wrong[:3] + unlisted_moved[:3])
full_audit("after Optimize+Apply (day 10)", ja, pa, sla, cfa, msa,
           book_store.load_last_applied_schedule(), frozen_keys_now())
ja2, pa2, *_ = run()
ADMIN.post("/optimize/done")
ja3, pa3, *_ = run()
check("repair right after Apply moves nothing", plan_hash(pa.schedule) == plan_hash(pa3.schedule),
      a=plan_hash(pa.schedule), b=plan_hash(pa3.schedule))

# =================================================================== PHASE 4
# Machine down (spec 5.8) on the day-10 state: waits, banner, Optimize, Apply.
use_store(fresh_dir(SNAP_DAY10, "phase-down"))
jd0, pd0, *_ = run()
start = date.fromisoformat(jd0["resolved_plan_start"])
d1 = next_working(start, cal)
d2 = next_working(d1, cal)
busy = defaultdict(int)
for e in pd0.schedule:
    if e.machine[:3] in ("CNC", "VMC") and any(d1 <= shift_day(s) <= d2 for s, _e, _w in e.op_segments or []):
        busy[e.machine] += 1
mach = max(busy, key=busy.get)
dates_before = dates_of(jd0)
rows_before = book_store.load_last_applied_schedule()
r = ADMIN.post("/machine-downtime", json={"machine": mach, "from_date": d1.isoformat(),
                                          "to_date": d2.isoformat(), "reason": "verify"})
assert r.status_code == 200, r.text
jd1, pd1, sld1, cfd1, msd1 = run()
ju1, *_ = run(USER)
a = full_audit(f"{mach} down {d1}..{d2}", jd1, pd1, sld1, cfd1, msd1, rows_before, frozen_keys_now())
on_m = sum(1 for e in pd1.schedule if e.machine == mach)
banner = [x for x in jd1["fixed_plan_alerts"] if mach in x]
check("machine down: banner names the machine, both roles",
      bool(banner) and ju1["fixed_plan_alerts"] == jd1["fixed_plan_alerts"],
      banner=banner[:1], ops_still_on_machine=on_m)
dm = moved(dates_before, dates_of(jd1))
check("machine down: published rows untouched by the break",
      book_store.load_last_applied_schedule() == rows_before, orders_whose_date_slid=len(dm))
m._start_optimize(ARGS.opt_budget, "deep", background=False)
st = ADMIN.get("/optimize/status").json()
dc = st.get("date_changes")
before = dates_of(run()[0])
ra = ADMIN.post("/optimize/apply")
jd2, pd2, *_ = run()
inc_d = m._incumbent_metrics()["total_late_days"]
check("machine down: the panel's After late-days is the plan after Apply (final review I2)",
      (st.get("best") or {}).get("total_late_days") == inc_d,
      panel=(st.get("best") or {}).get("total_late_days"), after_apply=inc_d)
ADMIN.post("/optimize/done")
jd3, pd3, *_ = run()
check("machine down: a repair right after Apply moves nothing (final review I3)",
      plan_hash(pd2.schedule) == plan_hash(pd3.schedule),
      a=plan_hash(pd2.schedule), b=plan_hash(pd3.schedule))
after = dates_of(jd2)
listed = {f"{r['so']}{KEY_SEP}{r['item']}": r["after"] for r in (dc or [])}
wrong = [k for k, v in listed.items() if after.get(k) != v]
unl = [k for k in moved(before, after) if k not in listed]
check("machine down: Optimize lists date changes, Apply clears the banner, list was true",
      dc is not None and ra.status_code == 200 and not [x for x in jd2["fixed_plan_alerts"] if mach in x]
      and not wrong and not unl,
      date_changes=None if dc is None else len(dc), alerts_after=jd2["fixed_plan_alerts"],
      wrong=len(wrong), unlisted_moved=len(unl),
      ops_on_machine_in_break_after=sum(1 for e in pd2.schedule if e.machine == mach and any(
          d1 <= shift_day(s) <= d2 for s, _e, _w in e.op_segments or [])))
RESULTS["machine_down"] = {"machine": mach, "from": d1.isoformat(), "to": d2.isoformat(),
                           "ops_queued_on_it": busy[mach], "orders_slid": len(dm)}

# =================================================================== PHASE 5
# Add New Orders (spec 5.7) on the day-0 published plan: quote -> add -> repairs.
CLOCK["now"] = datetime.fromisoformat(ARGS.start)
use_store(fresh_dir(SNAP_DAY0, "phase-quote"))
jq0, pq0, *_ = run()
d_before = dates_of(jq0)
mach_before = step_machines(pq0.schedule)
items = sorted({e.item_code for e in pq0.schedule}, key=lambda it: -sum(
    1 for e in pq0.schedule if e.item_code == it and e.machine[:3] in ("CNC", "VMC")))
drafts = [{"so_no": "VERIFY-NEW-1", "item_code": items[0], "qty": 40},
          {"so_no": "VERIFY-NEW-2", "item_code": items[1], "qty": 25}]
assert ADMIN.put("/new-orders/drafts", json={"drafts": drafts}).status_code == 200
q = ADMIN.post("/new-orders/quote").json()
quoted = {f"{l['so_no']}{KEY_SEP}{l['item_code']}": l["completion"] for l in q["lines"]}
ra = ADMIN.post("/new-orders/add", json={"stamp": q["stamp"]})
jq1, pq1, *_ = run()
d1_ = dates_of(jq1)
ex_moved = moved(d_before, d1_)
mach_after = step_machines(pq1.schedule)
mach_moved = [k for k, v in mach_before.items() if k in mach_after and mach_after[k] != v]
check("Add New Orders: quoted date = plan date, no existing order moved (date or machine)",
      q["verified"] and ra.status_code == 200 and not ra.json().get("warning")
      and all(d1_.get(k) == v for k, v in quoted.items()) and not ex_moved and not mach_moved,
      verified=q["verified"], quoted=list(quoted.values()),
      planned=[d1_.get(k) for k in quoted], existing_moved=len(ex_moved),
      machines_moved=len(mach_moved), warning=ra.json().get("warning"))
full_audit("after add", jq1, pq1, *run()[2:], book_store.load_last_applied_schedule(),
           frozen_keys_now())
ADMIN.post("/optimize/done")
jq2, pq2, *_ = run()
d2_ = dates_of(jq2)
check("Add New Orders: a repair right after keeps the quoted dates and every other date",
      all(d2_.get(k) == v for k, v in quoted.items()) and not moved(d1_, d2_),
      moved=len(moved(d1_, d2_)), planned=[d2_.get(k) for k in quoted])
# two punched days after the accept: record what the quoted dates do
Dq = date.fromisoformat(jq2["resolved_plan_start"])
qtrack = []
cur = pq2
for k in range(2):
    CLOCK["now"] = datetime.combine(Dq, datetime.min.time()) + timedelta(hours=23)
    for e, qq, _ends, who in day_plan(cur.schedule, Dq):
        punch(e, qq, Dq, who, ms0, ADMIN)
    ADMIN.post("/optimize/done")
    jx, cur, slx, cfx, msx = run()
    qtrack.append({"day": Dq.isoformat(), "quoted": [jx["expected_end"].get(k) for k in quoted]})
    full_audit(f"quote+punched day {k + 1}", jx, cur, slx, cfx, msx,
               book_store.load_last_applied_schedule(), frozen_keys_now())
    Dq = next_working(Dq, cal)
RESULTS["quote"] = {"quoted": list(quoted.values()), "track": qtrack}

# =================================================================== PHASE 6
# Add New Orders, "ask for an earlier date" (spec D10, final review C1) on the day-10
# state: preponed search -> accept (publishes) -> repair.
CLOCK["now"] = OPT_CLOCK
use_store(fresh_dir(SNAP_DAY10, "phase-prepone"))
m._OPT_BUDGETS["deep"] = ARGS.opt_budget
jp0, pp0, *_ = run()
items = sorted({e.item_code for e in pp0.schedule}, key=lambda it: -sum(
    1 for e in pp0.schedule if e.item_code == it and e.machine[:3] in ("CNC", "VMC")))
pdrafts = [{"so_no": "VERIFY-EARLY-1", "item_code": items[0], "qty": 40},
           {"so_no": "VERIFY-EARLY-2", "item_code": items[1], "qty": 25}]
assert ADMIN.put("/new-orders/drafts", json={"drafts": pdrafts}).status_code == 200
pq = ADMIN.post("/new-orders/quote").json()
targets = {f"{l['so_no']}{KEY_SEP}{l['item_code']}":
           (date.fromisoformat(l["completion"]) - timedelta(days=7)).isoformat()
           for l in pq["lines"]}
t = time.time()
rp = ADMIN.post("/new-orders/prepone", json={"targets": targets})
assert rp.status_code == 200, rp.text[:300]
while m._OPTIMIZE["state"] == "running" and time.time() - t < 3600:
    time.sleep(0.5)
pst = ADMIN.get("/optimize/status").json()
qm = pst.get("quote_movement") or {}
before = dates_of(run()[0])
racc = ADMIN.post("/new-orders/prepone/accept")
jp1, pp1, slp1, cfp1, msp1 = run()
after = dates_of(jp1)
new_sos = {d["so_no"] for d in pdrafts}
prow = book_store.load_last_applied_schedule()
pub_new = {(so, r["item_code"], r["process_seq"]): r["machine"] for r in prow
           for so in r["so_refs"] if so in new_sos}
plan_new = {k: v for k, v in step_machines(pp1.schedule).items() if k[0] in new_sos}
routed = {(d["so_no"], d["item_code"], op.seq) for d in pdrafts
          for op in msp1.routings[d["item_code"]].processes
          if (d["so_no"], d["item_code"], op.seq) in plan_new}
check("prepone accept: the published plan contains the accepted orders, and the plan holds "
      "them there (final review C1)",
      racc.status_code == 200 and routed and all(k in pub_new for k in routed)
      and all(plan_new[k] == {pub_new[k]} for k in routed),
      accept=racc.status_code, steps_planned=len(routed),
      steps_published=sum(1 for k in routed if k in pub_new),
      held=sum(1 for k in routed if k in pub_new and plan_new[k] == {pub_new[k]}))
told = {f"{r['so_no']}{KEY_SEP}{r['item_code']}": r["after"] for r in qm.get("moved") or []}
wrong = [k for k, v in told.items() if after.get(k) != v]
later_untold = [k for k, v in before.items() if k not in told and after.get(k, v) > v]
due = {f"{o.so_no}{KEY_SEP}{o.item_code}": o.delivery_date
       for o in book_store.load_active_orders().values()}
late_after = sum(max(0, (date.fromisoformat(after[k]) - due[k]).days)
                 for k in before if k in after and k in due)
check("prepone accept: the 'what moved' screen told the truth (final review C1/I2)",
      bool(qm) and not wrong and not later_untold and qm.get("late_days_after") == late_after,
      moved_listed=len(told), wrong=len(wrong), later_but_unlisted=len(later_untold),
      screen_late_after=qm.get("late_days_after"), floor_late_after=late_after,
      screen_late_before=qm.get("late_days_before"))
full_audit("after prepone accept", jp1, pp1, slp1, cfp1, msp1, prow, frozen_keys_now())
ADMIN.post("/optimize/done")
jp2, pp2, *_ = run()
check("prepone accept: a repair right after moves nothing",
      plan_hash(pp1.schedule) == plan_hash(pp2.schedule),
      a=plan_hash(pp1.schedule), b=plan_hash(pp2.schedule))
RESULTS["prepone"] = {"secs": round(time.time() - t, 1),
                      "new_orders_achieved": [after.get(k) for k in targets],
                      "targets": list(targets.values()),
                      "existing_moved_later": len(told),
                      "late_days_existing_before": qm.get("late_days_before"),
                      "late_days_existing_after": qm.get("late_days_after"),
                      "late_after_accept": late_days(pp1.schedule, slp1, cfp1)}

RESULTS["late_days"] = {"a_free_before_feature": late_a, "b_day0_repaired": late_b,
                        "c_day10_repaired": late_c, "c_day10_free_same_ranks": late_c_free,
                        "d_day10_optimize_search_best": RESULTS["optimize_day10"]["search_best_late"],
                        "d_day10_after_apply_repaired": late_d}
RESULTS["secs"] = round(time.time() - t0, 1)
npass = sum(1 for c in RESULTS["checks"] if c["pass"])
print(f"\nCHECKS {npass}/{len(RESULTS['checks'])} passed;  late-days {json.dumps(RESULTS['late_days'])}")
json.dump(RESULTS, open(ARGS.out, "w"), indent=1, default=str)
shutil.rmtree(WORK, ignore_errors=True)
